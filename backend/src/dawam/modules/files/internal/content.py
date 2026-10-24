"""Looking inside an upload: what it really is, and the text in it.

Nothing here trusts the client: the type comes from the bytes (magic numbers, the
container structure of DOCX/XLSX, strict UTF-8 for text) and must agree with the
file's extension; only types on the allow-list get through.
"""

from __future__ import annotations

import io
import logging
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from .ocr import OcrEngine
    from .parsing import DocumentParser

logger = logging.getLogger(__name__)

EXTRACTED_TEXT_MAX_CHARS = 2_000_000
PDF_MAX_PAGES = 1_000
DOCX_MAX_PARAGRAPHS = 50_000
XLSX_MAX_ROWS = 200_000
XLSX_MAX_CELLS = 2_000_000
"""Work caps: a huge (or blank) document must not keep a request busy for long."""
ZIP_MAX_ENTRIES = 5_000
ZIP_MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
"""A DOCX/XLSX that unpacks to more than this is refused (a zip bomb)."""


@dataclass(frozen=True)
class FileType:
    name: str
    mime: str
    extensions: tuple[str, ...]


PDF = FileType("pdf", "application/pdf", (".pdf",))
DOCX = FileType(
    "docx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    (".docx",),
)
XLSX = FileType(
    "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", (".xlsx",)
)
MARKDOWN = FileType("md", "text/markdown", (".md", ".markdown"))
TEXT = FileType("txt", "text/plain", (".txt",))
SQL = FileType("sql", "application/sql", (".sql",))
YAML = FileType("yaml", "application/yaml", (".yaml", ".yml"))
CSV = FileType("csv", "text/csv", (".csv",))
JSON = FileType("json", "application/json", (".json",))
PNG = FileType("png", "image/png", (".png",))
JPEG = FileType("jpeg", "image/jpeg", (".jpg", ".jpeg"))
GIF = FileType("gif", "image/gif", (".gif",))
WEBP = FileType("webp", "image/webp", (".webp",))

TEXT_TYPES = (MARKDOWN, TEXT, SQL, YAML, CSV, JSON)
"""Types stored as UTF-8 text: these can be edited in the browser (spec §6.17)."""
ALLOWED_TYPES = (PDF, DOCX, XLSX, *TEXT_TYPES, PNG, JPEG, GIF, WEBP)
"""The allow-list. SVG and HTML are deliberately absent: they can carry script."""

_IMAGES = (PNG, JPEG, GIF, WEBP)
_BY_EXTENSION = {ext: t for t in ALLOWED_TYPES for ext in t.extensions}


class Unsupported(Exception):
    """The file is not an allowed type, or is not what its extension says."""


def _zip_kind(data: bytes) -> FileType | None:
    if not data.startswith(b"PK"):
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > ZIP_MAX_ENTRIES:
                raise Unsupported("The file holds too many entries.")
            if sum(info.file_size for info in infos) > ZIP_MAX_UNCOMPRESSED_BYTES:
                raise Unsupported("The file unpacks to an unreasonable size.")
            names = {info.filename for info in infos}
    except zipfile.BadZipFile:
        return None
    if "[Content_Types].xml" not in names:
        return None
    if "word/document.xml" in names:
        return DOCX
    if "xl/workbook.xml" in names:
        return XLSX
    return None


def _binary_kind(data: bytes) -> FileType | None:
    if b"%PDF-" in data[:1024]:
        return PDF
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return PNG
    if data.startswith(b"\xff\xd8\xff"):
        return JPEG
    if data.startswith((b"GIF87a", b"GIF89a")):
        return GIF
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return WEBP
    return _zip_kind(data)


def _is_text(data: bytes) -> bool:
    if b"\x00" in data:
        return False
    try:
        data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return False
    return True


def detect(filename: str, data: bytes) -> FileType:
    """The type of ``data``, which must be on the allow-list and agree with the
    extension of ``filename``. Raises ``Unsupported``."""
    dot = filename.rfind(".")
    claimed = _BY_EXTENSION.get(filename[dot:].lower()) if dot >= 0 else None
    if claimed is None:
        raise Unsupported(
            "That file type is not allowed. Upload a PDF, DOCX, XLSX, Markdown, SQL, YAML, "
            "CSV, JSON, text or image (PNG, JPEG, GIF or WebP) file."
        )
    found = _binary_kind(data)
    if found is None and claimed in TEXT_TYPES and _is_text(data):
        found = claimed
    if found is None or found != claimed:
        raise Unsupported("The file's content does not match its extension.")
    return found


def extract_text(
    file_type: FileType,
    data: bytes,
    parser: DocumentParser | None = None,
    ocr: OcrEngine | None = None,
    on_ocr: Callable[[], None] | None = None,
) -> tuple[str | None, str]:
    """``(text, status)``: ``status`` is ``extracted``, ``no_text_found`` (e.g. a scanned
    PDF without ``ocr``, or a document that cannot be read) or ``none`` (images: nothing to
    extract).

    A ``parser`` that handles the type returns layout-aware Markdown; if it fails, the basic
    extraction (pypdf, python-docx, openpyxl) is used instead. With ``ocr``, a PDF that has
    pages short of text is read by OCR (after ``on_ocr``); if OCR fails or reads nothing,
    the text found without it stands."""
    if file_type in _IMAGES:
        return None, "none"
    try:
        if parser is not None and file_type in parser.types and _within_caps(file_type, data):
            try:
                text = parser.parse(file_type, data)
            except Exception:
                logger.warning("layout-aware parsing failed; using basic extraction", exc_info=True)
                text = _extract(file_type, data)
        else:
            text = _extract(file_type, data)
    except Exception:
        return None, "no_text_found"
    if ocr is not None and file_type is PDF:
        text = _with_ocr(ocr, data, text, on_ocr)
    text = text.replace("\x00", "").strip()[:EXTRACTED_TEXT_MAX_CHARS]
    if not text:
        return None, "no_text_found"
    return text, "extracted"


def _with_ocr(ocr: OcrEngine, data: bytes, text: str, on_ocr: Callable[[], None] | None) -> str:
    from .ocr import pages_needing_ocr

    try:
        pages = pages_needing_ocr(data)
    except Exception:
        return text
    if not pages:
        return text
    if on_ocr is not None:
        on_ocr()  # outside the try: a cancelled or superseded job stops here
    try:
        recognised = ocr.recognise(data, pages)
    except Exception:
        logger.warning("OCR failed; keeping the text found without it", exc_info=True)
        return text
    return recognised if recognised.strip() else text


def _within_caps(file_type: FileType, data: bytes) -> bool:
    """Whether a document is small enough for a parser that has no caps of its own; a bigger
    one is read by the basic extraction, which stops at the caps."""
    try:
        if file_type is PDF:
            from pypdf import PdfReader

            return len(PdfReader(io.BytesIO(data)).pages) <= PDF_MAX_PAGES
        if file_type is DOCX:
            from docx import Document

            document = Document(io.BytesIO(data))
            rows = sum(len(table.rows) for table in document.tables)
            return len(document.paragraphs) <= DOCX_MAX_PARAGRAPHS and rows <= DOCX_MAX_PARAGRAPHS
        if file_type is XLSX:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True)
            try:
                rows = cells = 0
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        rows += 1
                        cells += len(row)
                        if rows > XLSX_MAX_ROWS or cells > XLSX_MAX_CELLS:
                            return False
            finally:
                workbook.close()
    except Exception:
        return False  # unreadable here: let the basic extraction decide
    return True


class _Text:
    """Collects text pieces, refusing more once the character cap is reached."""

    def __init__(self) -> None:
        self._parts: list[str] = []
        self._size = 0

    @property
    def full(self) -> bool:
        return self._size >= EXTRACTED_TEXT_MAX_CHARS

    def add(self, piece: str) -> None:
        room = EXTRACTED_TEXT_MAX_CHARS - self._size
        if room <= 0:
            return
        piece = piece[:room]
        self._parts.append(piece)
        self._size += len(piece) + 1

    def text(self) -> str:
        return "\n".join(self._parts)


def _extract(file_type: FileType, data: bytes) -> str:
    if file_type in TEXT_TYPES:
        return data[: EXTRACTED_TEXT_MAX_CHARS * 4].decode("utf-8-sig", errors="ignore")
    out = _Text()
    if file_type is PDF:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            return ""
        for number, page in enumerate(reader.pages):
            if number >= PDF_MAX_PAGES or out.full:
                break
            out.add(page.extract_text() or "")
    elif file_type is DOCX:
        from docx import Document

        document = Document(io.BytesIO(data))
        for count, paragraph in enumerate(document.paragraphs):
            if count >= DOCX_MAX_PARAGRAPHS or out.full:
                break
            out.add(paragraph.text)
        rows = 0
        for table in document.tables:
            for row in table.rows:
                rows += 1
                if rows > DOCX_MAX_PARAGRAPHS or out.full:
                    break
                out.add("\t".join(cell.text for cell in row.cells))
    elif file_type is XLSX:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = cells = 0
        try:
            for sheet in workbook.worksheets:
                out.add(f"## {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    # Every row counts, blank or not: a huge empty sheet must end quickly.
                    rows += 1
                    cells += len(row)
                    if rows > XLSX_MAX_ROWS or cells > XLSX_MAX_CELLS or out.full:
                        break
                    line = "\t".join("" if cell is None else str(cell) for cell in row)
                    if line.strip("\t"):
                        out.add(line.rstrip("\t"))
                if rows > XLSX_MAX_ROWS or cells > XLSX_MAX_CELLS or out.full:
                    break
        finally:
            workbook.close()
    return out.text()
