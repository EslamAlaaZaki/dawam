"""The PII inventory of a Workspace (spec §6.12, story 140): the confirmed PII source
columns and the DW columns they feed, with category, hashing, handling note and reviewer, as
XLSX or Markdown for PDPL records of processing and audits.

PII-influenced tables (PII only steers them) are listed on their own, never among the PII
columns. Only metadata: no data value is ever read, so none can appear.
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from datetime import datetime
from typing import Literal

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from dawam.modules.auth import AuthService, User
from dawam.modules.sources import ConfirmedPiiColumn, confirmed_pii_columns
from dawam.modules.warehouse import pii_origins
from dawam.platform.clock import Clock

from .service import LineageService, PiiView

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MARKDOWN_MIME = "text/markdown; charset=utf-8"
InventoryFormat = Literal["xlsx", "md"]

SOURCE_SHEET = "Source columns"
DW_SHEET = "DW columns"
TABLE_SHEET = "Influenced tables"
SOURCE_HEADER = [
    "source_system",
    "database_schema",
    "table",
    "column",
    "category",
    "rules",
    "reviewer",
    "reviewed_at",
]
DW_HEADER = [
    "layer",
    "table",
    "column",
    "source_columns",
    "category",
    "hashing",
    "handling_note",
    "reviewer",
]
TABLE_HEADER = ["layer", "table"]
HASHED = "pseudonymised"
NOT_HASHED = "not hashed"

Section = tuple[str, list[str], list[list[str]]]


@dataclass(frozen=True)
class PiiInventory:
    name: str
    mime: str
    data: bytes


def _join(values: list[str]) -> str:
    return "; ".join(dict.fromkeys(v for v in values if v))


def _append(sheet: Worksheet, values: list[object]) -> None:
    """Append a row with every text cell stored as a string: a handling note starting with
    ``=``, ``+``, ``-`` or ``@`` must never become a formula."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _cell(value: str) -> str:
    """Text for a Markdown table cell: no pipes or line breaks to break the row."""
    return " ".join(value.replace("\\", "\\\\").replace("|", "\\|").split())


def _xlsx(sections: list[Section]) -> bytes:
    workbook = Workbook()
    for index, (title, header, rows) in enumerate(sections):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        assert sheet is not None
        sheet.title = title
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for values in rows:
            _append(sheet, list(values))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _markdown(sections: list[Section], generated: datetime) -> bytes:
    lines = ["# PII inventory", "", f"Generated {generated.date().isoformat()}.", ""]
    for title, header, rows in sections:
        lines += [f"## {title}", ""]
        if not rows:
            lines += ["None.", ""]
            continue
        lines += [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        lines += ["| " + " | ".join(_cell(v) for v in row) + " |" for row in rows]
        lines.append("")
    return "\n".join(lines).encode()


class PiiInventoryService:
    """Authorizes through the lineage service it wraps (any member reads)."""

    def __init__(
        self, engine: sa.Engine, lineage: LineageService, *, auth: AuthService, clock: Clock
    ) -> None:
        self._engine = engine
        self._lineage = lineage
        self._auth = auth
        self._clock = clock

    def export(
        self, user: User, workspace_id: uuid.UUID, *, fmt: InventoryFormat = "xlsx"
    ) -> PiiInventory:
        """The inventory as of now: source columns, PII-derived DW columns and, separately,
        PII-influenced DW tables."""
        view = self._lineage.pii_view(user, workspace_id)  # authorizes
        with Session(self._engine) as db:
            sources = confirmed_pii_columns(db, workspace_id)
            origins = pii_origins(db, [s.id for s in sources])
        reviewers = {s.reviewer_id for s in sources if s.reviewer_id}
        names = {i: u.display_name for i, u in self._auth.users_by_id(reviewers).items()}
        sections: list[Section] = [
            (SOURCE_SHEET, SOURCE_HEADER, self._source_rows(sources, names)),
            (DW_SHEET, DW_HEADER, self._dw_rows(view, sources, origins, names)),
            (TABLE_SHEET, TABLE_HEADER, [[t.layer, t.label] for t in view.tables]),
        ]
        if fmt == "md":
            data = _markdown(sections, self._clock())
            return PiiInventory("pii-inventory.md", MARKDOWN_MIME, data)
        return PiiInventory("pii-inventory.xlsx", XLSX_MIME, _xlsx(sections))

    @staticmethod
    def _source_rows(
        sources: list[ConfirmedPiiColumn], names: dict[uuid.UUID, str]
    ) -> list[list[str]]:
        return [
            [
                s.system_code,
                s.db_schema,
                s.table,
                s.column,
                s.category,
                _join(s.rules),
                names.get(s.reviewer_id, "") if s.reviewer_id else "",
                s.reviewed_at.date().isoformat() if s.reviewed_at else "",
            ]
            for s in sources
        ]

    @staticmethod
    def _dw_rows(
        view: PiiView,
        sources: list[ConfirmedPiiColumn],
        origins: dict[uuid.UUID, set[uuid.UUID]],
        names: dict[uuid.UUID, str],
    ) -> list[list[str]]:
        by_id = {s.id: s for s in sources}
        rows = []
        for column in view.columns:
            feeding = sorted(
                (by_id[i] for i in origins.get(column.id, ()) if i in by_id),
                key=lambda s: (s.system_code, s.db_schema, s.table, s.column),
            )
            table, _, name = column.label.rpartition(".")
            rows.append(
                [
                    column.layer,
                    table,
                    name,
                    _join([f"{s.system_code}.{s.db_schema}.{s.table}.{s.column}" for s in feeding]),
                    _join([s.category for s in feeding]),
                    HASHED if column.pseudonymised else NOT_HASHED,
                    column.pii_note or "",
                    _join([names.get(s.reviewer_id, "") if s.reviewer_id else "" for s in feeding]),
                ]
            )
        return rows
