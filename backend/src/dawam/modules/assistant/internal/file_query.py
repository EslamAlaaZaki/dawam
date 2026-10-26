"""The ``query_file`` engine (#398, ai-enhancements §6.2): one read-only SELECT over an
uploaded CSV or XLSX, run in-process by DuckDB in a fresh in-memory database per call.

In this order, failing closed at every step:

1. the file's bytes are written to a private temporary directory (an XLSX sheet as CSV) and
   loaded into a table; nothing else is ever written there, and it is deleted afterwards;
2. a column whose header passes a built-in PII name rule, or whose sampled values pass a PII
   validator (the test ``run_source_query`` masks by), becomes ``[masked]`` (NULL stays NULL,
   so it can still be counted) in the table ``data`` the model queries; the unmasked table
   is dropped, so no query can read such a value back, whole or piece by piece;
3. external access is then switched off and the configuration locked: SQL can reach no file,
   URL, extension or other database, and cannot switch it back on;
4. DuckDB's own parser must find exactly one SELECT statement: a write, a ``COPY``, ``SET``,
   ``PRAGMA``, ``ATTACH`` or a second statement is rejected before anything runs;
5. loading and the query run under one time limit (the connection is interrupted) and at
   most ``ROW_CAP`` rows are fetched;
6. result columns whose values pass a PII validator (a value the query built) are masked
   whole, and cells are cut like a source query's;
7. a failure is described in our own words: DuckDB's error text can echo a value.

Limits are code constants for now; they become admin settings with Document AI settings.
"""

from __future__ import annotations

import csv
import io
import tempfile
import threading
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import duckdb
import openpyxl

from dawam.modules.sources import MASK, match_name, personal_columns, shown_cell
from dawam.platform.errors import ApiError
from dawam.platform.pii_validators import redact_text

CSV_MIME = "text/csv"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ROW_CAP = 1000
TIME_LIMIT_SECONDS = 10.0
"""For loading the file and running the query together."""
MEMORY_LIMIT = "512MB"
SCAN_VALUES = 1000
"""Values of each column tested against the PII validators when the file is loaded."""
TABLE = "data"

_SETTINGS = {
    "autoinstall_known_extensions": False,
    "autoload_known_extensions": False,
    "allow_community_extensions": False,
    "threads": 1,
    "memory_limit": MEMORY_LIMIT,
}


@dataclass(frozen=True)
class FileColumn:
    name: str
    type: str
    masked: bool
    """Holds personal data: every value shows as ``[masked]``."""


@dataclass(frozen=True)
class FileQueryResult:
    columns: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    row_count: int
    truncated: bool
    """The query had more rows than the cap."""
    masked_columns: tuple[str, ...]
    """Result columns masked because their values passed a PII validator."""
    table: tuple[FileColumn, ...]
    """The columns of ``data``, to write the next query against."""
    duration_ms: int


def run_file_query(
    data: bytes,
    mime: str,
    sql: str,
    *,
    sheet: str | None = None,
    row_cap: int = ROW_CAP,
) -> FileQueryResult:
    """Run ``sql`` over the file as the table ``data``. 415 ``not_queryable``, 404
    ``sheet_not_found``, 422 ``query_rejected``, ``query_failed`` or ``query_timeout``."""
    if mime not in (CSV_MIME, XLSX_MIME):
        raise ApiError(415, "not_queryable", "Only CSV and XLSX files can be queried.")
    started = time.monotonic()
    con = duckdb.connect(":memory:", config=_SETTINGS)
    timer = threading.Timer(TIME_LIMIT_SECONDS, con.interrupt)
    timer.start()
    try:
        with tempfile.TemporaryDirectory(prefix="dawam-query-file-") as folder:
            path = Path(folder) / "data.csv"
            path.write_bytes(data if mime == CSV_MIME else _xlsx_as_csv(data, sheet))
            table = _load(con, path)
        con.execute("SET enable_external_access = false")
        con.execute("SET lock_configuration = true")
        statement = _one_select(con, sql)
        cursor = con.execute(statement)
        names = tuple(redact_text(str(d[0]), entities=False) for d in cursor.description or ())
        raw = cursor.fetchmany(row_cap + 1)
    except duckdb.InterruptException:
        raise ApiError(
            422, "query_timeout", f"The query did not finish within {TIME_LIMIT_SECONDS:g} s."
        ) from None
    except duckdb.Error as exc:
        raise _failed(exc) from None
    finally:
        timer.cancel()
        con.close()
    truncated = len(raw) > row_cap
    raw = raw[:row_cap]
    hits = personal_columns([[row[i] for row in raw] for i in range(len(names))])
    rows = tuple(
        tuple(MASK if i in hits else shown_cell(value) for i, value in enumerate(row))
        for row in raw
    )
    return FileQueryResult(
        columns=names,
        rows=rows,
        row_count=len(rows),
        truncated=truncated,
        masked_columns=tuple(names[i] for i in sorted(hits)),
        table=table,
        duration_ms=round((time.monotonic() - started) * 1000),
    )


def _xlsx_as_csv(data: bytes, sheet: str | None) -> bytes:
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (OSError, KeyError, ValueError, zipfile.BadZipFile):
        raise ApiError(422, "query_failed", "The file could not be read as a table.") from None
    try:
        if sheet is None:
            rows = book.worksheets[0].iter_rows(values_only=True)
        elif sheet in book.sheetnames:
            rows = book[sheet].iter_rows(values_only=True)
        else:
            raise ApiError(
                404,
                "sheet_not_found",
                "The workbook has no such sheet. Its sheets: " + ", ".join(book.sheetnames),
            )
        out = io.StringIO()
        csv.writer(out).writerows(("" if v is None else v for v in row) for row in rows)
        return out.getvalue().encode()
    finally:
        book.close()


def _quoted(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _load(con: duckdb.DuckDBPyConnection, path: Path) -> tuple[FileColumn, ...]:
    """Load the file into ``data`` with its personal-data columns masked."""
    try:
        con.execute(
            "CREATE TABLE raw AS SELECT * FROM read_csv(?, header = true, sample_size = -1)",
            [str(path)],
        )
    except duckdb.InterruptException:
        raise
    except duckdb.Error:
        raise ApiError(422, "query_failed", "The file could not be read as a table.") from None
    described = con.execute("DESCRIBE raw").fetchall()
    columns: list[FileColumn] = []
    selected: list[str] = []
    for name, kind, *_ in described:
        quoted = _quoted(name)
        sample = [
            row[0]
            for row in con.execute(
                f"SELECT {quoted} FROM raw WHERE {quoted} IS NOT NULL LIMIT {SCAN_VALUES}"
            ).fetchall()
        ]
        masked = match_name(name) is not None or bool(personal_columns([sample]))
        columns.append(FileColumn(name, "VARCHAR" if masked else kind, masked))
        selected.append(
            f"CASE WHEN {quoted} IS NULL THEN NULL ELSE '{MASK}' END AS {quoted}"
            if masked
            else quoted
        )
    con.execute(f"CREATE TABLE {TABLE} AS SELECT {', '.join(selected)} FROM raw")
    con.execute("DROP TABLE raw")
    return tuple(columns)


def _one_select(con: duckdb.DuckDBPyConnection, sql: str) -> duckdb.Statement:
    try:
        statements = con.extract_statements(sql)
    except duckdb.Error:
        raise ApiError(422, "query_rejected", "The SQL does not parse.") from None
    if len(statements) != 1 or statements[0].type != duckdb.StatementType.SELECT:
        raise ApiError(
            422, "query_rejected", f"Only one SELECT over the table `{TABLE}` can be run."
        )
    return statements[0]


_NAMED_ERRORS = (duckdb.ParserException, duckdb.BinderException, duckdb.CatalogException)


def _failed(exc: duckdb.Error) -> ApiError:
    """Parser, binder and catalog errors name the model's own SQL (redacted all the same);
    any other error text can quote the file, so it is replaced."""
    if isinstance(exc, _NAMED_ERRORS) and not isinstance(exc, duckdb.PermissionException):
        message = redact_text(str(exc).splitlines()[0])[:300]
    elif isinstance(exc, duckdb.PermissionException):
        message = f"Only the table `{TABLE}` can be read: no files, URLs or other databases."
    else:
        message = "DuckDB could not run that query on this file."
    return ApiError(422, "query_failed", message)
