"""The mapping sheet of a Data Warehouse (spec §6.14, story 108): every Core and Mart column's
mapping as XLSX or CSV in the familiar source-to-target layout, for people outside the tool."""

from __future__ import annotations

import csv
import io
import uuid
from dataclasses import dataclass
from typing import Literal

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from dawam.modules.auth import User
from dawam.modules.sources import StagingSourceService
from dawam.platform.errors import ApiError

from .ddl import translate_type
from .mapping_service import ColumnMappingView, MappingService, TableMappingView
from .model_service import ModelService
from .service import DataWarehouseService
from .tables import DwTableRecord

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv"
MAPPING_LAYERS = ("core", "mart")
ExportFormat = Literal["xlsx", "csv"]

MAPPING_HEADER = [
    "layer",
    "target_table",
    "branch",
    "target_column",
    "target_type",
    "input_layer",
    "input_system",
    "input_database_schema",
    "input_table",
    "input_column(s)",
    "transformation_rule",
    "sql_expression",
    "mapping_type",
    "lookup_dimension",
    "notes",
]
BRANCH_HEADER = [
    "layer",
    "target_table",
    "branch",
    "driving_input",
    "joins",
    "filters",
    "group_by",
    "having",
    "integration_rule",
    "match_keys",
]
BRANCH_SHEET = "Branches"
MAPPING_SHEET = "Mapping"


@dataclass(frozen=True)
class MappingSheet:
    name: str
    mime: str
    data: bytes


def _distinct(values: list[str]) -> str:
    return ";".join(dict.fromkeys(v for v in values if v))


def _append(sheet: Worksheet, values: list[object]) -> None:
    """Append a row with every text cell stored as a string: a rule or expression starting
    with ``=``, ``+``, ``-`` or ``@`` must never become a formula."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


class MappingExportService:
    """Authorizes through the services it wraps (any member reads)."""

    def __init__(
        self,
        engine: sa.Engine,
        warehouses: DataWarehouseService,
        model: ModelService,
        mappings: MappingService,
        sources: StagingSourceService | None = None,
    ) -> None:
        self._engine = engine
        self._warehouses = warehouses
        self._model = model
        self._mappings = mappings
        self._sources = sources or StagingSourceService(engine)

    def export(
        self,
        user: User,
        workspace_id: uuid.UUID,
        *,
        fmt: ExportFormat = "xlsx",
        layer: str | None = None,
    ) -> MappingSheet:
        """The sheet of one Layer, or of Core and Mart when ``layer`` is ``None``. CSV holds
        the column sheet only. 404 ``not_found`` before set up, 422 ``invalid_layer``."""
        if layer is not None and layer not in MAPPING_LAYERS:
            raise ApiError(422, "invalid_layer", "The Layer is core or mart.")
        warehouse = self._warehouses.get(user, workspace_id)
        if warehouse is None:
            raise ApiError(404, "not_found", "The Data Warehouse is not set up yet.")
        platform = warehouse.target_platform
        origins = self._staging_origins(workspace_id)
        mapping_rows: list[list[object]] = []
        branch_rows: list[list[object]] = []
        for summary in self._model.list_tables(user, workspace_id, layer=layer):
            if summary.layer not in MAPPING_LAYERS:
                continue
            table = self._model.get_table(user, workspace_id, summary.id)
            view = self._mappings.get_mapping(user, workspace_id, summary.id)
            types = {c.name: translate_type(platform, c.data_type) for c in table.columns}
            mapping_rows += self._column_rows(view, types, self._source_tables(view, origins))
            branch_rows += self._branch_rows(view)
        stem = f"mapping-sheet-{layer or 'data-warehouse'}"
        if fmt == "csv":
            return MappingSheet(f"{stem}.csv", CSV_MIME, self._csv(mapping_rows))
        return MappingSheet(f"{stem}.xlsx", XLSX_MIME, self._xlsx(mapping_rows, branch_rows))

    # --- rows ------------------------------------------------------------------------

    def _staging_origins(self, workspace_id: uuid.UUID) -> dict[uuid.UUID, tuple[str, str]]:
        """Each Source Table's ``(system code, database schema)``, by id."""
        return {
            t.id: (system.code, t.db_schema)
            for system in self._sources.read(workspace_id)
            for t in system.tables
        }

    def _source_tables(
        self, view: TableMappingView, origins: dict[uuid.UUID, tuple[str, str]]
    ) -> dict[uuid.UUID, tuple[str, str]]:
        """For each input table that is a Staging Table, where its data comes from."""
        ids = {
            i.table_id
            for c in [*view.columns, *(c for b in view.branches for c in b.columns)]
            for i in c.inputs
        }
        if not ids:
            return {}
        with Session(self._engine) as db:
            links = db.execute(
                sa.select(DwTableRecord.id, DwTableRecord.source_table_id).where(
                    DwTableRecord.id.in_(ids), DwTableRecord.source_table_id.is_not(None)
                )
            ).all()
        return {tid: origins[sid] for tid, sid in links if sid in origins}

    def _column_rows(
        self,
        view: TableMappingView,
        types: dict[str, str],
        origins: dict[uuid.UUID, tuple[str, str]],
    ) -> list[list[object]]:
        def row(branch: str, column: ColumnMappingView) -> list[object]:
            inputs = column.inputs
            return [
                view.layer,
                view.table_name,
                branch,
                column.column_name,
                types.get(column.column_name, ""),
                view.source_layer if inputs else "",
                _distinct([origins[i.table_id][0] for i in inputs if i.table_id in origins]),
                _distinct([origins[i.table_id][1] for i in inputs if i.table_id in origins]),
                _distinct([i.table_name for i in inputs]),
                ";".join(f"{i.table_name}.{i.column_name}" for i in inputs),
                column.rule_text,
                column.sql_expression,
                column.mapping_type,
                (column.lookup or {}).get("dimension_name") or "",
                view.notes,
            ]

        if not view.branches:
            return [row("", c) for c in view.columns]
        rows = [row(b.name, c) for b in view.branches for c in b.columns]
        rows += [row("", c) for c in view.columns if c.mapping_type == "system"]
        return rows

    @staticmethod
    def _branch_rows(view: TableMappingView) -> list[list[object]]:
        return [
            [
                view.layer,
                view.table_name,
                b.name,
                b.driving_input,
                b.joins,
                b.filters,
                b.group_by or "",
                b.having or "",
                view.integration_rule or "",
                ";".join(view.match_keys),
            ]
            for b in view.branches
        ]

    # --- formats ---------------------------------------------------------------------

    @staticmethod
    def _csv(rows: list[list[object]]) -> bytes:
        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\r\n")
        writer.writerow(MAPPING_HEADER)
        writer.writerows(rows)
        return b"\xef\xbb\xbf" + out.getvalue().encode()

    @staticmethod
    def _xlsx(mapping: list[list[object]], branches: list[list[object]]) -> bytes:
        workbook = Workbook()
        first = workbook.active
        assert first is not None
        first.title = MAPPING_SHEET
        second = workbook.create_sheet(BRANCH_SHEET)
        for sheet, header, rows in (
            (first, MAPPING_HEADER, mapping),
            (second, BRANCH_HEADER, branches),
        ):
            sheet.append(header)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
            for values in rows:
                _append(sheet, values)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
