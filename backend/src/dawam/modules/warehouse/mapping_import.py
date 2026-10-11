"""Mapping sheet import (spec §6.14, story 109): the sheet ``mapping_export`` writes, filled in
by people outside the tool, comes back as a validation report and, for what differs, one
Change Set (origin ``import``).

The import only proposes. It never creates a table, column or branch: a row naming one that
does not exist fails validation. Staging rows are ignored, as are the columns that are not
mappings (``target_type``, the ``input_*`` columns, ``notes``) and the integration rule and
match keys. A mapping is created or its type, rule and SQL changed; a ``lookup``, ``system``
or ``hashed`` mapping may appear only as it is. Approvals are not in the sheet and not touched
(accepting a change to a revoking field returns an approved Layer to draft, as any edit does).
Exporting and importing the same state therefore proposes nothing.
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass, field
from typing import Any

import sqlalchemy as sa
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from dawam.modules.auth import User
from dawam.modules.changesets import ChangeSetDetail, ChangeSetService, ProposedItem
from dawam.modules.workspaces import Action, WorkspaceService
from dawam.platform.errors import ApiError

from .lineage_sql import BRANCH_PARTS
from .mapping_export import BRANCH_SHEET, MAPPING_HEADER, MAPPING_LAYERS, MAPPING_SHEET
from .mapping_service import ColumnMappingView, MappingService, TableMappingView
from .propagation_handlers import ColumnMappingHandler, MappingBranchHandler
from .service import DataWarehouseService

SCOPE = {"kind": "mapping_import"}
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000
REQUIRED_COLUMNS = (
    "layer",
    "target_table",
    "branch",
    "target_column",
    "mapping_type",
    "transformation_rule",
    "sql_expression",
)
REQUIRED_BRANCH_COLUMNS = ("layer", "target_table", "branch", *BRANCH_PARTS)
TOOL_MADE_TYPES = ("lookup", "system", "hashed")
"""Types the tool makes itself; an import may only repeat them."""


@dataclass(frozen=True)
class ImportIssue:
    sheet: str
    row: int | None
    """The row in the sheet (the header is row 1)."""
    message: str


@dataclass
class MappingImportReport:
    errors: list[ImportIssue] = field(default_factory=list)
    """Rows that failed validation; they propose nothing."""
    rows: int = 0
    """Mapping rows and branch rows read."""
    ignored_staging_rows: int = 0
    unchanged: int = 0
    changes: int = 0
    """Items proposed (or that would be, on a dry run)."""


@dataclass(frozen=True)
class MappingImportResult:
    report: MappingImportReport
    change_set: ChangeSetDetail | None
    """``None`` on a dry run, and when nothing differs."""


def _invalid_file(message: str) -> ApiError:
    return ApiError(422, "invalid_sheet", message)


def _cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _read_sheet(
    workbook: Any, name: str, required: tuple[str, ...]
) -> list[tuple[int, dict[str, Any]]]:
    """The rows of one sheet as ``(row number, {header: cell})``, blank rows left out."""
    rows = workbook[name].iter_rows(values_only=True, max_row=MAX_ROWS + 1)
    header = next(rows, None)
    names = [_cell(h) for h in header or ()]
    missing = [c for c in required if c not in names]
    if missing:
        raise _invalid_file(f"The {name} sheet lacks the columns {', '.join(missing)}.")
    result = []
    for number, values in enumerate(rows, start=2):
        if number > MAX_ROWS:
            raise _invalid_file(f"The {name} sheet has more than {MAX_ROWS} rows.")
        if all(_cell(v) == "" for v in values):
            continue
        result.append((number, dict(zip(names, values, strict=False))))
    return result


class MappingImportService:
    """Authorizes as editing the DW Schema and mappings (owners and editors)."""

    def __init__(
        self,
        engine: sa.Engine,
        *,
        workspaces: WorkspaceService,
        warehouses: DataWarehouseService,
        mappings: MappingService,
        change_sets: ChangeSetService,
    ) -> None:
        self._engine = engine
        self._workspaces = workspaces
        self._warehouses = warehouses
        self._mappings = mappings
        self._change_sets = change_sets

    def import_sheet(
        self,
        user: User,
        workspace_id: uuid.UUID,
        data: bytes,
        *,
        filename: str = "",
        dry_run: bool = True,
    ) -> MappingImportResult:
        """Validate the XLSX; unless ``dry_run``, propose what differs as a Change Set. Rows
        that fail are reported and left out. 422 ``invalid_sheet`` for a file that is not a
        mapping sheet, 404 ``not_found`` before the Data Warehouse is set up."""
        self._workspaces.authorize(user, Action.EDIT_DW_SCHEMA, workspace_id)
        if self._warehouses.settings_of(workspace_id) is None:
            raise ApiError(404, "not_found", "The Data Warehouse is not set up yet.")
        if len(data) > MAX_BYTES:
            raise _invalid_file(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
        mapping_rows, branch_rows = self._read(data)
        report = MappingImportReport(rows=len(mapping_rows) + len(branch_rows))
        views = {
            (v.layer, v.table_name.lower()): v for v in self._mappings.table_views(workspace_id)
        }
        items: list[ProposedItem] = []
        with Session(self._engine) as db:
            self._column_items(db, workspace_id, views, mapping_rows, report, items)
            self._branch_items(db, workspace_id, views, branch_rows, report, items)
        report.changes = len(items)
        if dry_run or not items:
            return MappingImportResult(report, None)
        detail = self._change_sets.propose(
            user,
            workspace_id,
            origin="import",
            scope=SCOPE,
            title=f"Mapping sheet import{f' ({filename})' if filename else ''}"[:200],
            items=items,
        )
        return MappingImportResult(report, detail)

    # --- reading ---------------------------------------------------------------------

    @staticmethod
    def _read(data: bytes) -> tuple[list[tuple[int, dict]], list[tuple[int, dict]]]:
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise _invalid_file("The file is not a readable .xlsx workbook.") from None
        try:
            if MAPPING_SHEET not in workbook.sheetnames:
                raise _invalid_file(f"The workbook has no {MAPPING_SHEET} sheet.")
            if not set(MAPPING_HEADER) >= set(REQUIRED_COLUMNS):  # pragma: no cover
                raise AssertionError("the export lacks a column the import needs")
            mapping = _read_sheet(workbook, MAPPING_SHEET, REQUIRED_COLUMNS)
            branches = (
                _read_sheet(workbook, BRANCH_SHEET, REQUIRED_BRANCH_COLUMNS)
                if BRANCH_SHEET in workbook.sheetnames
                else []
            )
        finally:
            workbook.close()
        return mapping, branches

    # --- column mappings -------------------------------------------------------------

    def _column_items(
        self,
        db: Session,
        workspace_id: uuid.UUID,
        views: dict[tuple[str, str], TableMappingView],
        rows: list[tuple[int, dict]],
        report: MappingImportReport,
        items: list[ProposedItem],
    ) -> None:
        handler = ColumnMappingHandler()
        seen: set[tuple[str, str, str, str]] = set()

        def error(number: int, message: str) -> None:
            report.errors.append(ImportIssue(MAPPING_SHEET, number, message))

        for number, row in rows:
            layer = _cell(row["layer"]).lower()
            if layer == "staging":
                report.ignored_staging_rows += 1
                continue
            if layer not in MAPPING_LAYERS:
                error(number, f"The layer `{layer}` is not core, mart or staging.")
                continue
            table_name, branch_name = _cell(row["target_table"]), _cell(row["branch"])
            column_name = _cell(row["target_column"])
            view = views.get((layer, table_name.lower()))
            if view is None:
                error(number, f"There is no {layer} table `{table_name}`; imports create nothing.")
                continue
            columns = view.columns
            branch_id = None
            if branch_name:
                branch = next(
                    (b for b in view.branches if b.name.lower() == branch_name.lower()), None
                )
                if branch is None:
                    error(number, f"`{table_name}` has no branch `{branch_name}`.")
                    continue
                columns, branch_id = branch.columns, branch.id
            current = next(
                (c for c in columns if c.column_name.lower() == column_name.lower()), None
            )
            if current is None:
                if branch_name and any(
                    c.column_name.lower() == column_name.lower() for c in view.columns
                ):
                    error(number, f"`{column_name}` is mapped at table level, not in a branch.")
                else:
                    error(
                        number,
                        f"`{table_name}` has no column `{column_name}`; imports create nothing.",
                    )
                continue
            if view.branches and not branch_name and current.mapping_type != "system":
                error(number, f"`{table_name}` is mapped in branches; name the branch.")
                continue
            key = (layer, table_name.lower(), branch_name.lower(), column_name.lower())
            if key in seen:
                error(number, f"`{table_name}.{column_name}` appears twice.")
                continue
            seen.add(key)
            item = self._column_item(
                db, workspace_id, handler, number, row, view, current, branch_id, error
            )
            if item is False:
                continue
            if item is None:
                report.unchanged += 1
            else:
                items.append(item)

    def _column_item(
        self,
        db: Session,
        workspace_id: uuid.UUID,
        handler: ColumnMappingHandler,
        number: int,
        row: dict,
        view: TableMappingView,
        current: ColumnMappingView,
        branch_id: uuid.UUID | None,
        error: Any,
    ) -> ProposedItem | bool | None:
        """The item for one row, ``None`` when it matches, ``False`` when it failed."""
        mapping_type = _cell(row["mapping_type"]).lower()
        label = f"{view.layer}.{view.table_name}.{current.column_name}"
        if mapping_type in TOOL_MADE_TYPES or current.mapping_type in TOOL_MADE_TYPES:
            if mapping_type != current.mapping_type:
                error(
                    number,
                    f"{label} is a {current.mapping_type} mapping; lookup, system and hashed "
                    "mappings are changed in the tool, not by import.",
                )
                return False
            dimension = _cell((current.lookup or {}).get("dimension_name"))
            if mapping_type == "lookup" and _cell(row.get("lookup_dimension")) != dimension:
                error(number, f"{label} looks up `{dimension}`; change a lookup in the tool.")
                return False
            return None
        rule, sql = _cell(row["transformation_rule"]), _cell(row["sql_expression"])
        wanted = {"mapping_type": mapping_type, "rule_text": rule, "sql_expression": sql}
        have = {
            "mapping_type": current.mapping_type,
            "rule_text": current.rule_text.strip(),
            "sql_expression": current.sql_expression.strip(),
        }
        if wanted == have:
            return None
        try:
            handler.check(db, workspace_id, current.column_id, branch_id is not None, wanted)
        except ApiError as exc:
            error(number, f"{label}: {exc.message}")
            return False
        if current.id is None:
            return ProposedItem(
                key=f"create:{current.column_id}:{branch_id}",
                object_type="column_mapping",
                operation="create",
                label=label,
                payload={
                    "dw_column_id": str(current.column_id),
                    "branch_id": str(branch_id) if branch_id else None,
                    **wanted,
                },
            )
        changed = {f for f in wanted if wanted[f] != have[f]}
        if changed & {"mapping_type", "sql_expression"}:
            changed |= {"mapping_type", "sql_expression"}
        return ProposedItem(
            key=f"update:{current.id}",
            object_type="column_mapping",
            operation="update",
            object_id=current.id,
            label=label,
            payload={f: wanted[f] for f in wanted if f in changed},
        )

    # --- branches --------------------------------------------------------------------

    def _branch_items(
        self,
        db: Session,
        workspace_id: uuid.UUID,
        views: dict[tuple[str, str], TableMappingView],
        rows: list[tuple[int, dict]],
        report: MappingImportReport,
        items: list[ProposedItem],
    ) -> None:
        handler = MappingBranchHandler()
        seen: set[tuple[str, str, str]] = set()

        def error(number: int, message: str) -> None:
            report.errors.append(ImportIssue(BRANCH_SHEET, number, message))

        for number, row in rows:
            layer = _cell(row["layer"]).lower()
            if layer == "staging":
                report.ignored_staging_rows += 1
                continue
            table_name, branch_name = _cell(row["target_table"]), _cell(row["branch"])
            view = views.get((layer, table_name.lower()))
            if view is None:
                error(number, f"There is no {layer} table `{table_name}`; imports create nothing.")
                continue
            branch = next((b for b in view.branches if b.name.lower() == branch_name.lower()), None)
            if branch is None:
                error(
                    number, f"`{table_name}` has no branch `{branch_name}`; imports create nothing."
                )
                continue
            key = (layer, table_name.lower(), branch_name.lower())
            if key in seen:
                error(number, f"Branch `{branch_name}` of `{table_name}` appears twice.")
                continue
            seen.add(key)
            wanted: dict[str, str | None] = {
                part: _cell(row[part]) or (None if part in ("group_by", "having") else "")
                for part in BRANCH_PARTS
            }
            have = {
                part: getattr(branch, part) or (None if part in ("group_by", "having") else "")
                for part in BRANCH_PARTS
            }
            changed = {p: v for p, v in wanted.items() if v != have[p]}
            if not changed:
                report.unchanged += 1
                continue
            label = f"{view.layer}.{view.table_name}: {branch.name}"
            try:
                handler.check(db, workspace_id, view.table_id, {"name": branch.name, **wanted})
            except ApiError as exc:
                error(number, f"{label}: {exc.message}")
                continue
            items.append(
                ProposedItem(
                    key=f"branch:{branch.id}",
                    object_type="mapping_branch",
                    operation="update",
                    object_id=branch.id,
                    label=label,
                    payload=changed,
                )
            )
