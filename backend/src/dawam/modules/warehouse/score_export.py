"""The score report of a Data Warehouse as Markdown or XLSX (spec §6.11, story 124): the
overall, per-Layer and per-table scores with every failed check and its fix hint, for
design reviews."""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from dawam.modules.auth import User

from .score_service import FailedCheck, Score, ScoreService

MARKDOWN_MIME = "text/markdown; charset=utf-8"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ScoreReportFormat = Literal["md", "xlsx"]

SUMMARY_SHEET = "Summary"
LAYERS_SHEET = "Layers"
TABLES_SHEET = "Tables"
CHECKS_SHEET = "Failed checks"
SUMMARY_HEADER = ["measure", "value"]
LAYERS_HEADER = ["layer", "score", "grade", "tables"]
TABLES_HEADER = ["layer", "table", "score"]
CHECKS_HEADER = [
    "severity",
    "layer",
    "check",
    "title",
    "category",
    "object_type",
    "object",
    "message",
    "fix_hint",
    "link",
]


@dataclass(frozen=True)
class ScoreReportFile:
    name: str
    mime: str
    data: bytes


def _number(value: float | None) -> str:
    return "not scored" if value is None else f"{value:g}"


def _cell(text: str) -> str:
    """Text for one Markdown line or table cell: no pipes, no line breaks."""
    return re.sub(r"\s+", " ", text).replace("|", "\\|").strip()


def _append(sheet: Worksheet, values: list[object]) -> None:
    """Append a row with every text cell stored as a string, so a message or name starting
    with ``=``, ``+``, ``-`` or ``@`` never becomes a formula."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


class ScoreExportService:
    """Authorizes through ``ScoreService`` (any member reads)."""

    def __init__(self, scores: ScoreService) -> None:
        self._scores = scores

    def export(
        self, user: User, workspace_id: uuid.UUID, *, fmt: ScoreReportFormat = "md"
    ) -> ScoreReportFile:
        """The latest score as a report. 404 ``not_found`` before the Data Warehouse is set up."""
        score = self._scores.current(user, workspace_id)
        if fmt == "xlsx":
            return ScoreReportFile("score-report.xlsx", XLSX_MIME, self._xlsx(score))
        return ScoreReportFile("score-report.md", MARKDOWN_MIME, self._markdown(score).encode())

    # --- formats ---------------------------------------------------------------------

    @staticmethod
    def _markdown(score: Score) -> str:
        lines = [
            "# Data Warehouse score report",
            "",
            f"Calculated at {score.calculated_at.isoformat()}.",
            "",
            "## Overall",
            "",
            f"- Score: {_number(score.score)}",
            f"- Grade: {score.grade or 'not scored'}"
            + (" (capped at C by an unresolved error)" if score.capped else ""),
            f"- Checks passed: {score.checks_passed}",
            f"- Checks failed: {score.checks_failed}",
            "",
            "## Layers",
            "",
            "| Layer | Score | Grade | Tables |",
            "| --- | --- | --- | --- |",
        ]
        lines += [
            f"| {item.layer} | {_number(item.score)} | {item.grade or '-'} | {item.table_count} |"
            for item in score.layers
        ]
        lines += ["", "## Tables", ""]
        if score.tables:
            lines += ["| Layer | Table | Score |", "| --- | --- | --- |"]
            lines += [f"| {t.layer} | {_cell(t.name)} | {t.score:g} |" for t in score.tables]
        else:
            lines.append("No tables are scored yet.")
        lines += ["", "## Failed checks", ""]
        if not score.failed_checks:
            lines.append("No failed checks.")
        for check in score.failed_checks:
            lines += ScoreExportService._markdown_check(check)
        return "\n".join(lines).rstrip() + "\n"

    @staticmethod
    def _markdown_check(check: FailedCheck) -> list[str]:
        return [
            f"### {check.severity.upper()}: {_cell(check.title)} ({check.layer})",
            "",
            f"- Object: {check.object_type} `{_cell(check.object_name)}` ({check.check_code})",
            f"- Problem: {_cell(check.message)}",
            f"- Fix: {_cell(check.fix_hint)}",
            "",
        ]

    @staticmethod
    def _xlsx(score: Score) -> bytes:
        workbook = Workbook()
        summary = workbook.active
        assert summary is not None
        summary.title = SUMMARY_SHEET
        sheets: list[tuple[Worksheet, list[str], list[list[object]]]] = [
            (
                summary,
                SUMMARY_HEADER,
                [
                    ["score", score.score],
                    ["grade", score.grade],
                    ["capped", score.capped],
                    ["calculated_at", score.calculated_at.isoformat()],
                    ["checks_passed", score.checks_passed],
                    ["checks_failed", score.checks_failed],
                ],
            ),
            (
                workbook.create_sheet(LAYERS_SHEET),
                LAYERS_HEADER,
                [[i.layer, i.score, i.grade, i.table_count] for i in score.layers],
            ),
            (
                workbook.create_sheet(TABLES_SHEET),
                TABLES_HEADER,
                [[t.layer, t.name, t.score] for t in score.tables],
            ),
            (
                workbook.create_sheet(CHECKS_SHEET),
                CHECKS_HEADER,
                [
                    [
                        c.severity,
                        c.layer,
                        c.check_code,
                        c.title,
                        c.category,
                        c.object_type,
                        c.object_name,
                        c.message,
                        c.fix_hint,
                        c.link,
                    ]
                    for c in score.failed_checks
                ],
            ),
        ]
        for sheet, header, rows in sheets:
            sheet.append(header)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
            for values in rows:
                _append(sheet, values)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
