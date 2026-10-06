"""The data dictionary workbook of a Source System (spec §6.15, story 129)."""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Font

from dawam.modules.auth import User

from .service import SourceSystemService
from .snapshot_service import SnapshotService, SourceSchema

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

TABLE_HEADER = [
    "schema",
    "table",
    "kind",
    "status",
    "description",
    "tags",
    "classification",
    "scd_hint",
    "is_sensitive",
    "source_comment",
]
COLUMN_HEADER = [
    "schema",
    "table",
    "column",
    "position",
    "data_type",
    "nullable",
    "primary_key",
    "default",
    "status",
    "description",
    "tags",
    "is_sensitive",
    "pii_category",
    "source_comment",
]


@dataclass(frozen=True)
class DataDictionary:
    name: str
    mime: str
    data: bytes


def _yes(value: bool | None) -> str:
    return "yes" if value else "no"


def build_dictionary(schema: SourceSchema) -> bytes:
    """A ``Tables`` sheet and a ``Columns`` sheet of the latest Snapshot's metadata with the
    descriptions, tags, classifications and PII categories users added. Only metadata:
    no row values are ever read, so none can appear."""
    workbook = Workbook()
    tables = workbook.active
    assert tables is not None
    tables.title = "Tables"
    columns = workbook.create_sheet("Columns")
    tables.append(TABLE_HEADER)
    columns.append(COLUMN_HEADER)
    for sheet in (tables, columns):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    for table in schema.content.tables:
        tables.append(
            [
                table.db_schema,
                table.name,
                table.kind,
                table.status,
                table.description,
                ", ".join(table.tags or []),
                table.classification,
                table.scd_hint,
                _yes(table.is_sensitive),
                table.comment,
            ]
        )
        for column in table.columns:
            columns.append(
                [
                    table.db_schema,
                    table.name,
                    column.name,
                    column.ordinal,
                    column.data_type,
                    _yes(column.is_nullable),
                    _yes(column.is_pk),
                    column.default,
                    column.status,
                    column.description,
                    ", ".join(column.tags or []),
                    _yes(column.is_sensitive),
                    column.pii_category,
                    column.comment,
                ]
            )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class DataDictionaryService:
    """Authorizes through the snapshot and system services it wraps (any member reads)."""

    def __init__(self, snapshots: SnapshotService, systems: SourceSystemService) -> None:
        self._snapshots = snapshots
        self._systems = systems

    def export(self, user: User, workspace_id: uuid.UUID, system_id: uuid.UUID) -> DataDictionary:
        """The workbook for the latest Snapshot. 404 before the first Snapshot."""
        system = self._systems.get(user, workspace_id, system_id)
        schema = self._snapshots.source_schema(user, workspace_id, system_id)
        return DataDictionary(
            name=f"data-dictionary-{system.code}.xlsx",
            mime=XLSX_MIME,
            data=build_dictionary(schema),
        )
