"""The Schema Import template, and reading an upload of it into a catalog (spec §6.4).

``build_template`` makes the Excel workbook an Editor downloads and fills: one sheet per
metadata kind plus a ``README`` sheet that documents every column. ``read_catalog`` turns
an upload (the filled workbook, or one CSV per sheet) into the ``SourceCatalog`` the
Snapshot writer takes, together with a validation report. Nothing here touches the
database: the caller saves the catalog only when the report allows it.

A report has errors (the import cannot be saved) and warnings (it can, once the user
accepts them). A file that cannot be read at all is an error in the report, never an
exception.
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook

from .connector import (
    ColumnInfo,
    ConstraintInfo,
    IndexInfo,
    RoutineInfo,
    SourceCatalog,
    TableInfo,
)

README_SHEET = "readme"
MAX_ROWS_PER_SHEET = 500_000
MAX_ISSUES = 200
"""A report stops listing issues past this many, so a hopeless file stays readable."""


@dataclass(frozen=True)
class TemplateColumn:
    name: str
    meaning: str
    required: bool = True


@dataclass(frozen=True)
class TemplateSheet:
    name: str
    purpose: str
    columns: tuple[TemplateColumn, ...]
    required: bool = True
    """Without it the import cannot be read (``tables`` and ``columns``)."""

    @property
    def column_names(self) -> tuple[str, ...]:
        return tuple(c.name for c in self.columns)


_LIST_NOTE = " Several columns: separate their names with commas."

SHEETS: tuple[TemplateSheet, ...] = (
    TemplateSheet(
        "schemas",
        "The Database Schemas to import, empty ones included. Optional: schemas of the "
        "other sheets count anyway.",
        (TemplateColumn("schema", "Database Schema name."),),
        required=False,
    ),
    TemplateSheet(
        "tables",
        "Tables and views.",
        (
            TemplateColumn("schema", "Database Schema holding the table or view."),
            TemplateColumn("table", "Table or view name."),
            TemplateColumn("kind", "`table` or `view`."),
            TemplateColumn("row_count", "Estimated number of rows (optional).", required=False),
            TemplateColumn("comment", "The database comment on it (optional).", required=False),
            TemplateColumn("definition", "A view's SQL (views only).", required=False),
        ),
    ),
    TemplateSheet(
        "columns",
        "Columns of the tables and views.",
        (
            TemplateColumn("schema", "Database Schema of the table."),
            TemplateColumn("table", "Table or view name."),
            TemplateColumn("column", "Column name."),
            TemplateColumn("data_type", "The type as the database spells it, e.g. `varchar(40)`."),
            TemplateColumn(
                "ordinal",
                "Position in the table, 1 first. Leave empty to keep the row order.",
                required=False,
            ),
            TemplateColumn("is_nullable", "`true` or `false`; empty means `true`.", required=False),
            TemplateColumn("default", "The default expression (optional).", required=False),
            TemplateColumn("comment", "The database comment on it (optional).", required=False),
        ),
    ),
    TemplateSheet(
        "constraints",
        "Primary keys, foreign keys and unique constraints.",
        (
            TemplateColumn("schema", "Database Schema of the table."),
            TemplateColumn("table", "The table the constraint belongs to."),
            TemplateColumn("constraint", "Constraint name."),
            TemplateColumn("type", "`pk`, `fk` or `unique`."),
            TemplateColumn("columns", "The constrained columns." + _LIST_NOTE),
            TemplateColumn("ref_schema", "Foreign keys: the referenced schema.", required=False),
            TemplateColumn("ref_table", "Foreign keys: the referenced table.", required=False),
            TemplateColumn(
                "ref_columns",
                "Foreign keys: the referenced columns, in the same order." + _LIST_NOTE,
                required=False,
            ),
        ),
        required=False,
    ),
    TemplateSheet(
        "indexes",
        "Indexes.",
        (
            TemplateColumn("schema", "Database Schema of the table."),
            TemplateColumn("table", "The table the index belongs to."),
            TemplateColumn("index", "Index name."),
            TemplateColumn(
                "columns", "The indexed columns or expressions." + _LIST_NOTE, required=True
            ),
            TemplateColumn("is_unique", "`true` or `false`; empty means `false`.", required=False),
        ),
        required=False,
    ),
    TemplateSheet(
        "routines",
        "Stored procedures and functions, with their source code.",
        (
            TemplateColumn("schema", "Database Schema holding the routine."),
            TemplateColumn("name", "Routine name."),
            TemplateColumn("kind", "`procedure` or `function`."),
            TemplateColumn(
                "signature",
                "The argument list that tells overloads apart (empty if the engine has none).",
                required=False,
            ),
            TemplateColumn("definition", "The routine's source code (optional).", required=False),
        ),
        required=False,
    ),
)
_SHEET_BY_NAME = {sheet.name: sheet for sheet in SHEETS}

_TRUE = {"true", "yes", "y", "1", "t"}
_FALSE = {"false", "no", "n", "0", "f"}
_CONSTRAINT_TYPES = {
    "pk": "pk",
    "primary key": "pk",
    "primary_key": "pk",
    "fk": "fk",
    "foreign key": "fk",
    "foreign_key": "fk",
    "unique": "unique",
    "uq": "unique",
}


def build_template() -> bytes:
    """The Schema Import workbook: a ``README`` sheet and one empty sheet per kind."""
    workbook = Workbook()
    readme = workbook.active
    assert readme is not None
    readme.title = "README"
    readme.append(["sheet", "column", "required", "meaning"])
    for sheet in SHEETS:
        readme.append([sheet.name, "", "required" if sheet.required else "optional", sheet.purpose])
        for column in sheet.columns:
            readme.append(
                [
                    sheet.name,
                    column.name,
                    "required" if column.required else "optional",
                    column.meaning,
                ]
            )
    readme.append([])
    readme.append(
        [
            "Fill each sheet below the header row and upload the workbook; each sheet can also "
            "be uploaded as a CSV named like the sheet (e.g. `columns.csv`). Nothing is saved "
            "until the upload has been checked."
        ]
    )
    for sheet in SHEETS:
        workbook.create_sheet(sheet.name).append(list(sheet.column_names))
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


@dataclass(frozen=True)
class Issue:
    sheet: str | None
    row: int | None
    """The row number in the sheet (the header is row 1); ``None`` for a whole sheet or file."""
    message: str


@dataclass
class ImportReport:
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)
    schema_count: int = 0
    table_count: int = 0
    column_count: int = 0
    routine_count: int = 0

    def error(self, sheet: str | None, row: int | None, message: str) -> None:
        if len(self.errors) < MAX_ISSUES:
            self.errors.append(Issue(sheet, row, message))

    def warn(self, sheet: str | None, row: int | None, message: str) -> None:
        if len(self.warnings) < MAX_ISSUES:
            self.warnings.append(Issue(sheet, row, message))

    @property
    def clean(self) -> bool:
        return not self.errors and not self.warnings


@dataclass(frozen=True)
class _Row:
    number: int
    cells: dict[str, str | None]


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _sheet_rows(
    name: str, rows: Iterable[Iterable[object]], report: ImportReport
) -> list[_Row] | None:
    """Header-keyed rows of one sheet; ``None`` (and an error) if the header is wrong."""
    iterator = iter(rows)
    header_cells = next(iterator, None)
    spec = _SHEET_BY_NAME[name]
    if header_cells is None:
        report.error(name, None, f"The sheet `{name}` is empty: it needs a header row.")
        return None
    header = [(_text(cell) or "").lower() for cell in header_cells]
    missing = [c.name for c in spec.columns if c.required and c.name not in header]
    if missing:
        report.error(name, 1, "Required column(s) missing: " + ", ".join(missing) + ".")
        return None
    extra = [h for h in header if h and h not in spec.column_names]
    if extra:
        report.warn(name, 1, "Unknown column(s) ignored: " + ", ".join(extra) + ".")
    seen: set[str] = set()
    for column in header:
        if column and column in seen:
            report.error(name, 1, f"The column `{column}` appears twice.")
            return None
        seen.add(column)
    result: list[_Row] = []
    for number, cells in enumerate(iterator, start=2):
        if number - 1 > MAX_ROWS_PER_SHEET:
            report.error(name, None, f"The sheet `{name}` has more than {MAX_ROWS_PER_SHEET} rows.")
            return None
        values = [_text(cell) for cell in cells]
        if not any(values):
            continue
        values += [None] * (len(header) - len(values))
        result.append(
            _Row(
                number,
                {h: v for h, v in zip(header, values, strict=False) if h in spec.column_names},
            )
        )
    return result


def _read_workbook(content: bytes, report: ImportReport) -> dict[str, list[_Row]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError, TypeError):
        report.error(None, None, "The file is not a readable Excel workbook (.xlsx).")
        return {}
    sheets: dict[str, list[_Row]] = {}
    try:
        for worksheet in workbook.worksheets:
            name = worksheet.title.strip().lower()
            if name == README_SHEET:
                continue
            if name not in _SHEET_BY_NAME:
                report.warn(
                    None,
                    None,
                    f"The sheet `{worksheet.title}` is not part of the template: ignored.",
                )
                continue
            rows = _sheet_rows(name, worksheet.iter_rows(values_only=True), report)
            if rows is not None:
                sheets[name] = rows
    except (zipfile.BadZipFile, KeyError, ValueError, OSError, TypeError):
        report.error(None, None, "The workbook is damaged and could not be read.")
        return {}
    finally:
        workbook.close()
    return sheets


def _read_csv(name: str, content: bytes, report: ImportReport) -> list[_Row] | None:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        report.error(name, None, f"The CSV for `{name}` is not UTF-8 text.")
        return None
    try:
        return _sheet_rows(name, csv.reader(io.StringIO(text)), report)
    except csv.Error:
        report.error(name, None, f"The CSV for `{name}` is malformed.")
        return None


def read_sheets(files: list[tuple[str, bytes]], report: ImportReport) -> dict[str, list[_Row]]:
    """The sheets in the uploaded files: a workbook gives all its sheets, a CSV the sheet
    its file name says (``columns.csv``)."""
    sheets: dict[str, list[_Row]] = {}

    def add(name: str, rows: list[_Row], source: str) -> None:
        if name in sheets:
            report.error(name, None, f"The sheet `{name}` is given twice (again in {source}).")
        else:
            sheets[name] = rows

    for filename, content in files:
        lowered = filename.lower()
        if lowered.endswith(".csv"):
            stem = lowered.rsplit("/", 1)[-1].rsplit("\\", 1)[-1][: -len(".csv")]
            if stem not in _SHEET_BY_NAME:
                report.error(
                    None,
                    None,
                    f"`{filename}` is not named after a sheet "
                    f"({', '.join(s.name for s in SHEETS)}).",
                )
                continue
            rows = _read_csv(stem, content, report)
            if rows is not None:
                add(stem, rows, filename)
        elif lowered.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
            for name, rows in _read_workbook(content, report).items():
                add(name, rows, filename)
        else:
            report.error(None, None, f"`{filename}` is neither an .xlsx workbook nor a .csv file.")
    return sheets


def _split(value: str | None) -> tuple[str, ...]:
    return tuple(part.strip() for part in value.split(",") if part.strip()) if value else ()


def _flag(report: ImportReport, sheet: str, row: _Row, column: str, *, default: bool) -> bool:
    value = row.cells.get(column)
    if value is None:
        return default
    lowered = value.lower()
    if lowered in _TRUE:
        return True
    if lowered in _FALSE:
        return False
    report.error(sheet, row.number, f"`{column}` must be true or false, not `{value}`.")
    return default


def _whole(report: ImportReport, sheet: str, row: _Row, column: str, *, minimum: int) -> int | None:
    value = row.cells.get(column)
    if value is None:
        return None
    try:
        number = int(value)
    except ValueError:
        try:
            as_float = float(value)
        except ValueError:
            report.error(sheet, row.number, f"`{column}` must be a whole number, not `{value}`.")
            return None
        if not as_float.is_integer():
            report.error(sheet, row.number, f"`{column}` must be a whole number, not `{value}`.")
            return None
        number = int(as_float)
    if number < minimum:
        report.error(sheet, row.number, f"`{column}` must be at least {minimum}.")
        return None
    return number


def _required(report: ImportReport, sheet: str, row: _Row, *columns: str) -> bool:
    missing = [c for c in columns if row.cells.get(c) is None]
    for column in missing:
        report.error(sheet, row.number, f"`{column}` is empty.")
    return not missing


def read_catalog(files: list[tuple[str, bytes]]) -> tuple[SourceCatalog | None, ImportReport]:
    """Parse and validate an upload. The catalog is ``None`` if the report has errors."""
    report = ImportReport()
    sheets = read_sheets(files, report)
    for spec in SHEETS:
        if (
            spec.required
            and spec.name not in sheets
            and not any(e.sheet == spec.name for e in report.errors)
        ):
            report.error(spec.name, None, f"The sheet `{spec.name}` is missing.")
    if report.errors:
        return None, report

    # -- tables -------------------------------------------------------------------------
    tables: dict[tuple[str, str], dict] = {}
    for row in sheets["tables"]:
        if not _required(report, "tables", row, "schema", "table", "kind"):
            continue
        schema, name = row.cells["schema"], row.cells["table"]
        kind = (row.cells["kind"] or "").lower()
        if kind not in ("table", "view"):
            report.error("tables", row.number, f"`kind` must be table or view, not `{kind}`.")
            continue
        assert schema is not None and name is not None
        if (schema, name) in tables:
            report.error("tables", row.number, f"The table `{schema}.{name}` is listed twice.")
            continue
        definition = row.cells.get("definition")
        if kind == "view" and definition is None:
            report.warn("tables", row.number, f"The view `{schema}.{name}` has no definition.")
        if kind == "table" and definition is not None:
            report.warn(
                "tables", row.number, f"The table `{schema}.{name}` has a definition: ignored."
            )
            definition = None
        tables[(schema, name)] = {
            "kind": kind,
            "row_estimate": _whole(report, "tables", row, "row_count", minimum=0),
            "comment": row.cells.get("comment"),
            "definition": definition,
            "columns": [],
            "constraints": [],
            "indexes": [],
        }

    # -- columns ------------------------------------------------------------------------
    for row in sheets["columns"]:
        if not _required(report, "columns", row, "schema", "table", "column", "data_type"):
            continue
        key = (row.cells["schema"] or "", row.cells["table"] or "")
        table = tables.get(key)
        if table is None:
            report.error(
                "columns", row.number, f"The table `{key[0]}.{key[1]}` is not in `tables`."
            )
            continue
        name = row.cells["column"] or ""
        if any(c["name"] == name for c in table["columns"]):
            report.error("columns", row.number, f"The column `{key[1]}.{name}` is listed twice.")
            continue
        table["columns"].append(
            {
                "name": name,
                "ordinal": _whole(report, "columns", row, "ordinal", minimum=1),
                "data_type": row.cells["data_type"],
                "is_nullable": _flag(report, "columns", row, "is_nullable", default=True),
                "default": row.cells.get("default"),
                "comment": row.cells.get("comment"),
                "row": row.number,
            }
        )
    for (schema, name), table in tables.items():
        if not table["columns"]:
            report.warn("columns", None, f"`{schema}.{name}` has no columns.")

    # -- constraints --------------------------------------------------------------------
    for row in sheets.get("constraints", []):
        if not _required(
            report, "constraints", row, "schema", "table", "constraint", "type", "columns"
        ):
            continue
        key = (row.cells["schema"] or "", row.cells["table"] or "")
        table = tables.get(key)
        if table is None:
            report.error(
                "constraints", row.number, f"The table `{key[0]}.{key[1]}` is not in `tables`."
            )
            continue
        kind = _CONSTRAINT_TYPES.get((row.cells["type"] or "").lower())
        if kind is None:
            report.error(
                "constraints",
                row.number,
                f"`type` must be pk, fk or unique, not `{row.cells['type']}`.",
            )
            continue
        name = row.cells["constraint"] or ""
        if any(c.name == name for c in table["constraints"]):
            report.error(
                "constraints",
                row.number,
                f"The constraint `{name}` is listed twice for `{key[1]}`.",
            )
            continue
        columns = _split(row.cells["columns"])
        known = {c["name"] for c in table["columns"]}
        unknown = [c for c in columns if c not in known]
        if unknown:
            report.error(
                "constraints",
                row.number,
                f"`{name}` names columns `{key[1]}` does not have: {', '.join(unknown)}.",
            )
            continue
        ref_schema = row.cells.get("ref_schema")
        ref_table = row.cells.get("ref_table")
        ref_columns = _split(row.cells.get("ref_columns"))
        if kind == "fk":
            if ref_table is None or not ref_columns:
                report.error(
                    "constraints",
                    row.number,
                    f"The foreign key `{name}` needs ref_table and ref_columns.",
                )
                continue
            if len(ref_columns) != len(columns):
                report.error(
                    "constraints",
                    row.number,
                    f"The foreign key `{name}` has {len(columns)} column(s) but "
                    f"{len(ref_columns)} referenced.",
                )
                continue
            ref_schema = ref_schema or key[0]
            target = tables.get((ref_schema, ref_table))
            if target is None:
                report.warn(
                    "constraints",
                    row.number,
                    f"The foreign key `{name}` points at `{ref_schema}.{ref_table}`, which is not "
                    "in this import.",
                )
            else:
                target_columns = {c["name"] for c in target["columns"]}
                absent = [c for c in ref_columns if c not in target_columns]
                if absent:
                    report.error(
                        "constraints",
                        row.number,
                        f"The foreign key `{name}` points at columns `{ref_table}` does not have: "
                        f"{', '.join(absent)}.",
                    )
                    continue
        else:
            ref_schema, ref_table, ref_columns = None, None, ()
        table["constraints"].append(
            ConstraintInfo(
                name=name,
                type=kind,
                columns=columns,
                ref_schema=ref_schema,
                ref_table=ref_table,
                ref_columns=ref_columns,
            )
        )

    # -- indexes ------------------------------------------------------------------------
    for row in sheets.get("indexes", []):
        if not _required(report, "indexes", row, "schema", "table", "index", "columns"):
            continue
        key = (row.cells["schema"] or "", row.cells["table"] or "")
        table = tables.get(key)
        if table is None:
            report.error(
                "indexes", row.number, f"The table `{key[0]}.{key[1]}` is not in `tables`."
            )
            continue
        name = row.cells["index"] or ""
        if any(i.name == name for i in table["indexes"]):
            report.error(
                "indexes", row.number, f"The index `{name}` is listed twice for `{key[1]}`."
            )
            continue
        table["indexes"].append(
            IndexInfo(
                name=name,
                columns=_split(row.cells["columns"]),
                is_unique=_flag(report, "indexes", row, "is_unique", default=False),
            )
        )

    # -- routines -----------------------------------------------------------------------
    routines: list[RoutineInfo] = []
    seen_routines: set[tuple[str, str, str, str]] = set()
    for row in sheets.get("routines", []):
        if not _required(report, "routines", row, "schema", "name", "kind"):
            continue
        kind = (row.cells["kind"] or "").lower()
        if kind not in ("procedure", "function"):
            report.error(
                "routines", row.number, f"`kind` must be procedure or function, not `{kind}`."
            )
            continue
        schema, name = row.cells["schema"] or "", row.cells["name"] or ""
        signature = row.cells.get("signature") or ""
        identity = (schema, kind, signature, name)
        if identity in seen_routines:
            report.error(
                "routines",
                row.number,
                f"The routine `{schema}.{name}({signature})` is listed twice.",
            )
            continue
        seen_routines.add(identity)
        routines.append(RoutineInfo(schema, name, kind, row.cells.get("definition"), signature))

    # -- schemas ------------------------------------------------------------------------
    listed = [r.cells["schema"] for r in sheets.get("schemas", []) if r.cells.get("schema")]
    used = {s for s, _ in tables} | {r.schema for r in routines}
    schema_names = list(dict.fromkeys(s for s in listed if s))
    if listed:
        for name in sorted(used - set(schema_names)):
            report.warn(
                "schemas", None, f"The schema `{name}` is used but not listed in `schemas`."
            )
    schema_names += sorted(used - set(schema_names))

    if report.errors:
        return None, report

    # -- the catalog --------------------------------------------------------------------
    table_infos: list[TableInfo] = []
    for (schema, name), table in tables.items():
        pk_columns = {c for k in table["constraints"] if k.type == "pk" for c in k.columns}
        keep_order = sorted(
            enumerate(table["columns"]),
            key=lambda item: (item[1]["ordinal"] is None, item[1]["ordinal"] or 0, item[0]),
        )
        columns = tuple(
            ColumnInfo(
                name=c["name"],
                ordinal=c["ordinal"] if c["ordinal"] is not None else position,
                data_type=c["data_type"],
                is_nullable=c["is_nullable"],
                is_pk=c["name"] in pk_columns,
                default=c["default"],
                comment=c["comment"],
            )
            for position, (_, c) in enumerate(keep_order, start=1)
        )
        table_infos.append(
            TableInfo(
                schema=schema,
                name=name,
                kind=table["kind"],
                row_estimate=table["row_estimate"],
                comment=table["comment"],
                definition=table["definition"],
                columns=columns,
                constraints=tuple(table["constraints"]),
                indexes=tuple(table["indexes"]),
            )
        )
    report.schema_count = len(schema_names)
    report.table_count = len(table_infos)
    report.column_count = sum(len(t.columns) for t in table_infos)
    report.routine_count = len(routines)
    return (
        SourceCatalog(
            tables=tuple(table_infos), routines=tuple(routines), schemas=tuple(schema_names)
        ),
        report,
    )
