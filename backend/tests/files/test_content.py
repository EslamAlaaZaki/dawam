"""Work caps of text extraction: huge or blank documents end quickly."""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook

from dawam.modules.files.internal import content


def xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def docx(paragraphs: list[str]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_blank_rows_count_toward_the_row_cap(monkeypatch):
    monkeypatch.setattr(content, "XLSX_MAX_ROWS", 10)
    data = xlsx([[None, None]] * 50 + [["findme"]])

    text, status = content.extract_text(content.XLSX, data)

    assert status == "extracted"  # only the sheet heading
    assert "findme" not in text


def test_rows_within_the_cap_are_read(monkeypatch):
    monkeypatch.setattr(content, "XLSX_MAX_ROWS", 10)

    text, status = content.extract_text(content.XLSX, xlsx([["a", "b"], [None], ["c"]]))

    assert status == "extracted"
    assert "a\tb" in text
    assert "c" in text


def test_the_cell_cap_stops_a_wide_sheet(monkeypatch):
    monkeypatch.setattr(content, "XLSX_MAX_CELLS", 10)

    text, _ = content.extract_text(content.XLSX, xlsx([["x"] * 6, ["y"] * 6, ["z"] * 6]))

    assert "z" not in text


def test_text_stops_being_built_at_the_character_cap():
    data = docx([f"paragraph {n:04d} " + "x" * 20 for n in range(1000)])

    text, status = content.extract_text(content.DOCX, data, max_chars=50)

    assert status == "extracted"
    assert len(text) <= 50
    assert "paragraph 0999" not in text


def test_the_paragraph_cap_limits_a_docx(monkeypatch):
    monkeypatch.setattr(content, "DOCX_MAX_PARAGRAPHS", 3)

    text, _ = content.extract_text(content.DOCX, docx([f"para{n}" for n in range(10)]))

    assert text == "para0\npara1\npara2"


# -- the Workspace's PDF page cap (Document AI ``max_pdf_pages``) ----------------------


class _Parser:
    types = (content.PDF,)

    def __init__(self) -> None:
        self.seen = 0

    def parse(self, file_type, data: bytes) -> str:
        self.seen += 1
        return "# Parsed\n\nby the layout parser"


class _Ocr:
    def __init__(self) -> None:
        self.pages: list[list[int]] = []

    def recognise(self, data: bytes, pages) -> str:
        self.pages.append(list(pages))
        return "read by OCR"


def three_pages() -> bytes:
    from tests.files.test_files import text_pdf
    from tests.files.test_ocr import pdf_of

    return pdf_of(text_pdf("one"), text_pdf("two"), text_pdf("three"))


def test_the_page_cap_decides_whether_the_parser_reads_a_pdf():
    parser = _Parser()

    content.extract_text(content.PDF, three_pages(), parser, max_pages=2)
    assert parser.seen == 0  # over the cap: the basic extraction reads it, up to the cap

    content.extract_text(content.PDF, three_pages(), parser, max_pages=3)
    assert parser.seen == 1


def test_the_page_cap_decides_whether_a_pdf_goes_to_ocr():
    ocr = _Ocr()

    content.extract_text(content.PDF, three_pages(), None, ocr, max_pages=2)
    assert ocr.pages == []

    content.extract_text(content.PDF, three_pages(), None, ocr, max_pages=5)
    assert ocr.pages == [[1, 2, 3]]  # each page holds a few characters: short of text


def test_the_engines_page_ceiling_is_the_settings_hard_limit():
    from dawam.modules.llm.document_ai import BY_KEY

    assert BY_KEY["max_pdf_pages"].high == content.PDF_PAGE_CEILING
