"""File storage and document upload (spec §6.17, stories 64 and 67, §4.3)."""

from __future__ import annotations

import io
import re
import uuid
from datetime import timedelta
from pathlib import Path

import pytest
import sqlalchemy as sa
from docx import Document
from openpyxl import Workbook

from tests.roles import RoleClients

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


def text_pdf(text: str) -> bytes:
    """A one-page PDF with ``text`` on it."""
    stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def blank_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Customer master data owner")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Status"
    table.rows[0].cells[1].text = "Dormant"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Codes"
    sheet.append(["code", "meaning"])
    sheet.append(["D", "Dormant account"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def system(roles: RoleClients) -> str:
    response = roles.client("owner").post(
        f"/api/v1/workspaces/{roles.workspace_id}/systems", json={"name": "Core", "code": "cbs"}
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def files_path(roles: RoleClients, system_id: str) -> str:
    return f"/api/v1/workspaces/{roles.workspace_id}/systems/{system_id}/files"


def upload(roles: RoleClients, system_id: str, name: str, data: bytes, *, as_role="editor"):
    return roles.client(as_role).post(
        files_path(roles, system_id),
        files={"file": (name, data, "application/octet-stream")},
    )


def error_code(response) -> str:
    return response.json()["error"]["code"]


def stored(app, file_id: str):
    """The row's extracted text, status and storage key. A deliberate storage-property
    check: the API does not expose extracted text until search exists."""
    with app.state.engine.connect() as conn:
        return conn.execute(
            sa.text(
                "SELECT extracted_text, text_status, storage_key FROM workspace_files "
                "WHERE id = :id"
            ),
            {"id": file_id},
        ).one()


def other_workspace(roles: RoleClients) -> dict:
    response = roles.client("non_member").post(
        "/api/v1/workspaces", json={"name": "Other", "description": "", "domain": ""}
    )
    assert response.status_code == 201, response.text
    return response.json()


@pytest.mark.parametrize(
    ("name", "data", "mime", "needle"),
    [
        ("sad.md", b"# Architecture\nThe ledger", "text/markdown", "The ledger"),
        ("notes.txt", "مرحبا".encode(), "text/plain", "مرحبا"),
        ("spec.pdf", text_pdf("Hello DAWAM"), "application/pdf", "Hello DAWAM"),
        (
            "mdm.docx",
            docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Dormant",
        ),
        (
            "codes.xlsx",
            xlsx_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "Dormant account",
        ),
    ],
    ids=["md", "txt", "pdf", "docx", "xlsx"],
)
def test_an_editor_uploads_a_document_and_its_text_is_extracted(
    roles: RoleClients, name, data, mime, needle
):
    system_id = system(roles)

    response = upload(roles, system_id, name, data)

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["name"] == name
    assert body["mime"] == mime
    assert body["size"] == len(data)
    assert body["kind"] == "uploaded"
    assert body["owner_kind"] == "source_system"
    assert body["owner_id"] == system_id
    assert body["text_status"] == "extracted"
    assert body["updated_by"] == str(roles.user("editor").id)
    text, status, _ = stored(roles.app, body["id"])
    assert status == "extracted"
    assert needle in text


def test_a_scanned_pdf_is_marked_no_text_found(roles: RoleClients):
    response = upload(roles, system(roles), "scan.pdf", blank_pdf())

    assert response.status_code == 201
    assert response.json()["text_status"] == "no_text_found"
    assert stored(roles.app, response.json()["id"])[0] is None


def test_an_image_is_stored_without_extraction(roles: RoleClients):
    response = upload(roles, system(roles), "erd.png", PNG)

    assert response.status_code == 201
    assert response.json()["mime"] == "image/png"
    assert response.json()["text_status"] == "none"


def test_any_member_downloads_a_file_as_an_attachment(roles: RoleClients):
    system_id = system(roles)
    file = upload(roles, system_id, "Plan ü.md", b"# Plan").json()
    url = f"/api/v1/workspaces/{roles.workspace_id}/files/{file['id']}/download"

    for role in ("viewer", "editor", "owner"):
        response = roles.client(role).get(url)
        assert response.status_code == 200
        assert response.content == b"# Plan"
        assert response.headers["content-type"].startswith("text/markdown")
        disposition = response.headers["content-disposition"]
        assert disposition.startswith("attachment;")
        assert "filename*=UTF-8''Plan%20%C3%BC.md" in disposition
        assert response.headers["x-content-type-options"] == "nosniff"


def test_non_members_and_anonymous_cannot_download(roles: RoleClients):
    file = upload(roles, system(roles), "a.md", b"secret").json()
    url = f"/api/v1/workspaces/{roles.workspace_id}/files/{file['id']}/download"

    assert roles.client("non_member").get(url).status_code == 404
    assert roles.client("admin").get(url).status_code == 404
    assert roles.client("anonymous").get(url).status_code == 401


def test_a_file_is_not_reachable_through_another_workspace(roles: RoleClients):
    file = upload(roles, system(roles), "a.md", b"secret").json()
    other = other_workspace(roles)

    # The non-member owns "Other", so only the file's own Workspace may serve it.
    response = roles.client("non_member").get(
        f"/api/v1/workspaces/{other['id']}/files/{file['id']}/download"
    )

    assert response.status_code == 404
    assert error_code(response) == "not_found"
    unknown = roles.client("owner").get(
        f"/api/v1/workspaces/{roles.workspace_id}/files/{uuid.uuid4()}/download"
    )
    assert unknown.status_code == 404


def test_a_system_of_another_workspace_is_not_found(roles: RoleClients):
    other = other_workspace(roles)
    foreign = roles.client("non_member").post(
        f"/api/v1/workspaces/{other['id']}/systems", json={"name": "X", "code": "x"}
    )
    assert foreign.status_code == 201

    response = upload(roles, foreign.json()["id"], "a.md", b"hi", as_role="owner")

    assert response.status_code == 404


def test_a_viewer_cannot_upload(roles: RoleClients):
    response = upload(roles, system(roles), "a.md", b"hi", as_role="viewer")

    assert response.status_code == 403


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("evil.pdf", b"MZ\x90\x00 not a pdf"),
        ("evil.png", b"<html><script>alert(1)</script></html>"),
        ("page.html", b"<html></html>"),
        ("logo.svg", b"<svg onload=alert(1)></svg>"),
        ("run.exe", b"MZ\x90\x00"),
        ("noextension", b"hello"),
        ("renamed.docx", xlsx_bytes()),
        ("renamed.md", PNG),
        ("fake.docx", b"PK\x03\x04 not really a zip"),
        ("binary.txt", b"\x00\x01\x02\xff\xfe"),
    ],
    ids=lambda value: value if isinstance(value, str) else "",
)
def test_content_that_is_not_an_allowed_type_is_refused(roles: RoleClients, name, data):
    response = upload(roles, system(roles), name, data)

    assert response.status_code == 415
    assert error_code(response) == "unsupported_file_type"


def test_a_file_over_the_limit_is_refused(roles: RoleClients):
    system_id = system(roles)
    roles.app.state.settings.upload_max_mb = 0.0001  # about 100 bytes

    response = upload(roles, system_id, "big.md", b"x" * 200)

    assert response.status_code == 413
    assert error_code(response) == "file_too_large"
    assert roles.client("owner").get(files_path(roles, system_id)).json()["items"] == []


def test_the_limit_is_25_mb_by_default(roles: RoleClients):
    assert roles.app.state.settings.upload_max_bytes == 25 * 1024 * 1024


def test_a_path_in_the_file_name_never_reaches_the_disk(roles: RoleClients, tmp_path: Path):
    system_id = system(roles)

    response = upload(roles, system_id, "..\\..\\evil/../../etc/passwd.md", b"x")

    assert response.status_code == 201
    assert response.json()["name"] == "passwd.md"
    paths = [p for p in tmp_path.rglob("*") if p.is_file()]
    assert len(paths) == 1
    assert tmp_path / "files" in paths[0].parents
    assert re.fullmatch(r"[0-9a-f]{32}", paths[0].name)  # a random key, not the name
    assert stored(roles.app, response.json()["id"])[2] == paths[0].name


def test_uploading_the_same_name_overwrites_in_place(roles: RoleClients, tmp_path: Path):
    system_id = system(roles)
    first = upload(roles, system_id, "a.md", b"one").json()

    second = upload(roles, system_id, "a.md", b"two two").json()

    assert second["id"] == first["id"]
    assert second["size"] == 7
    listed = roles.client("viewer").get(files_path(roles, system_id)).json()["items"]
    assert [f["name"] for f in listed] == ["a.md"]
    url = f"/api/v1/workspaces/{roles.workspace_id}/files/{first['id']}/download"
    assert roles.client("viewer").get(url).content == b"two two"
    assert len([p for p in tmp_path.rglob("*") if p.is_file()]) == 1  # the old bytes are gone


def test_the_file_area_lists_files_by_name_with_pages(roles: RoleClients):
    system_id = system(roles)
    for name in ("c.md", "a.md", "b.md"):
        assert upload(roles, system_id, name, b"x").status_code == 201

    viewer = roles.client("viewer")
    page = viewer.get(files_path(roles, system_id), params={"limit": 2}).json()
    assert [f["name"] for f in page["items"]] == ["a.md", "b.md"]
    rest = viewer.get(
        files_path(roles, system_id), params={"limit": 2, "cursor": page["next_cursor"]}
    ).json()
    assert [f["name"] for f in rest["items"]] == ["c.md"]
    assert rest["next_cursor"] is None


def test_uploads_are_recorded_in_the_activity_feed(roles: RoleClients, clock):
    system_id = system(roles)
    upload(roles, system_id, "a.md", b"one")
    clock.advance(timedelta(minutes=1))
    upload(roles, system_id, "a.md", b"two")

    feed = roles.client("owner").get(f"/api/v1/workspaces/{roles.workspace_id}/activity").json()

    events = [(i["verb"], i["object_label"]) for i in feed["items"] if i["object_type"] == "file"]
    assert events == [("file.replaced", "a.md"), ("file.uploaded", "a.md")]
