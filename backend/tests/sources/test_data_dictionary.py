"""Data dictionary export (spec §6.15, story 129): any member downloads a Source System's
data dictionary as XLSX, and an owner or editor saves it into the file area.

Behaviour is driven through the HTTP API; the workbook is read back with openpyxl.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from tests.roles import RoleClients
from tests.sample_source import SampleSource
from tests.sources.test_extraction import add_system, connect, extract
from tests.sources.test_pii_findings import by_column, findings
from tests.sources.test_schema_browser import schema

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def enhanced(roles: RoleClients, sample_source: SampleSource) -> str:
    """A Source System with a Snapshot, a described, tagged, classified table and a column
    confirmed as PII."""
    system = add_system(roles)
    connect(roles, system, sample_source.connection_body())
    extract(roles, system)
    editor = roles.client("editor")
    [table] = [
        t
        for t in schema(roles, system)["tables"]
        if (t["db_schema"], t["name"]) == ("core", "customers")
    ]
    response = editor.patch(
        f"{system}/tables/{table['id']}",
        json={
            "version": table["version"],
            "description": "Bank customers",
            "tags": ["master-data"],
            "classification": "master",
        },
    )
    assert response.status_code == 200, response.text
    finding = by_column(findings(roles, system))[("customers", "national_id")]
    assert editor.post(f"{system}/pii-findings/{finding['id']}/confirm").status_code == 200
    return system


def sheets(data: bytes) -> dict[str, list[tuple]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in workbook.worksheets}


def test_a_member_downloads_the_dictionary_with_tables_and_columns(
    roles: RoleClients, sample_source: SampleSource
):
    system = enhanced(roles, sample_source)

    response = roles.client("viewer").get(f"{system}/data-dictionary")

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    assert "data-dictionary-cbs.xlsx" in response.headers["content-disposition"]
    book = sheets(response.content)
    assert set(book) == {"Tables", "Columns"}
    tables = book["Tables"]
    header = list(tables[0])
    [customers] = [
        r for r in tables[1:] if (r[header.index("schema")], r[1]) == ("core", "customers")
    ]
    assert customers[header.index("description")] == "Bank customers"
    assert customers[header.index("tags")] == "master-data"
    assert customers[header.index("classification")] == "master"
    columns = book["Columns"]
    cheader = list(columns[0])
    [national_id] = [
        r
        for r in columns[1:]
        if r[cheader.index("table")] == "customers" and r[cheader.index("column")] == "national_id"
    ]
    assert national_id[cheader.index("data_type")]
    assert national_id[cheader.index("pii_category")]
    assert national_id[cheader.index("is_sensitive")] == "yes"


def test_the_dictionary_never_carries_values(roles: RoleClients, sample_source: SampleSource):
    system = enhanced(roles, sample_source)

    cells = {
        str(c)
        for rows in sheets(roles.client("viewer").get(f"{system}/data-dictionary").content).values()
        for row in rows
        for c in row
    }

    assert not any("@" in c and "." in c for c in cells), "an email-like value leaked"


def test_the_dictionary_needs_a_snapshot(roles: RoleClients):
    system = add_system(roles)

    assert roles.client("viewer").get(f"{system}/data-dictionary").status_code == 404


def test_an_editor_saves_the_dictionary_into_the_file_area(
    roles: RoleClients, sample_source: SampleSource
):
    system = enhanced(roles, sample_source)

    saved = roles.client("editor").post(f"{system}/files/data-dictionary")

    assert saved.status_code == 201, saved.text
    assert saved.json()["name"] == "data-dictionary-cbs.xlsx"
    assert saved.json()["owner_kind"] == "source_system"
    listed = roles.client("viewer").get(f"{system}/files").json()["items"]
    assert [f["name"] for f in listed] == ["data-dictionary-cbs.xlsx"]
    again = roles.client("editor").post(f"{system}/files/data-dictionary")
    assert again.status_code == 201 and again.json()["id"] == saved.json()["id"]


def test_a_viewer_cannot_save_the_dictionary(roles: RoleClients, sample_source: SampleSource):
    system = enhanced(roles, sample_source)

    assert roles.client("viewer").post(f"{system}/files/data-dictionary").status_code == 403


def test_text_starting_like_a_formula_is_stored_as_text(
    roles: RoleClients, sample_source: SampleSource
):
    system = enhanced(roles, sample_source)
    [table] = [
        t
        for t in schema(roles, system)["tables"]
        if (t["db_schema"], t["name"]) == ("core", "customers")
    ]
    formula = '=HYPERLINK("http://evil.example","click")'
    response = roles.client("editor").patch(
        f"{system}/tables/{table['id']}",
        json={"version": table["version"], "description": formula},
    )
    assert response.status_code == 200, response.text

    data = roles.client("viewer").get(f"{system}/data-dictionary").content

    sheet = load_workbook(io.BytesIO(data))["Tables"]
    [cell] = [c for row in sheet.iter_rows() for c in row if c.value == formula]
    assert cell.data_type == "s"
