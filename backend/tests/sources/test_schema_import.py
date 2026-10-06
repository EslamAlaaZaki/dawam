"""Schema Import: template, upload and validation (spec §6.4, stories 47, 49-51).

Behaviour is driven through the HTTP API. Uploads are CSVs (one per sheet), which the
same reader handles as workbook sheets; the workbook path is covered in
``test_schema_import_parsing``.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from tests.roles import RoleClients
from tests.sources import test_extraction
from tests.sources.test_extraction import (
    add_system,
    connect,
    extract,
    latest,
    snapshots,
    tables_of,
)

scratch_source = test_extraction.scratch_source

TABLES = "schema,table,kind,row_count\nsales,Customer,table,5\n"
COLUMNS = (
    "schema,table,column,data_type\nsales,Customer,id,integer\nsales,Customer,name,varchar(40)\n"
)
CONSTRAINTS = "schema,table,constraint,type,columns\nsales,Customer,pk,pk,id\n"


def csvs(tables: str = TABLES, columns: str = COLUMNS, **others: str) -> list[tuple]:
    sheets = {"tables": tables, "columns": columns, **others}
    return [("files", (f"{name}.csv", text.encode(), "text/csv")) for name, text in sheets.items()]


def upload(roles: RoleClients, system: str, files, as_role="editor", **form) -> dict:
    response = roles.client(as_role).post(
        f"{system}/import/upload",
        files=files,
        data={k: str(v).lower() if isinstance(v, bool) else v for k, v in form.items()},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_the_template_is_a_workbook_with_a_sheet_per_kind(roles: RoleClients):
    system = add_system(roles)

    response = roles.client("editor").get(f"{system}/import/template")

    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    names = load_workbook(io.BytesIO(response.content)).sheetnames
    assert names == ["README", "schemas", "tables", "columns", "constraints", "indexes", "routines"]


def test_an_upload_becomes_a_snapshot_with_origin_import(roles: RoleClients):
    system = add_system(roles)

    outcome = upload(roles, system, csvs(constraints=CONSTRAINTS))

    assert outcome["imported"] and not outcome["unchanged"]
    assert outcome["report"]["errors"] == [] and outcome["report"]["warnings"] == []
    assert outcome["snapshot"]["origin"] == "import"
    content = latest(roles, system)
    table = tables_of(content)[("sales", "Customer")]
    assert table["row_estimate"] == 5
    assert [(c["name"], c["is_pk"]) for c in table["columns"]] == [("id", True), ("name", False)]


def test_a_malformed_upload_writes_nothing(roles: RoleClients):
    system = add_system(roles)
    broken = csvs(columns="schema,table,column\nsales,Customer,id\n")

    outcome = upload(roles, system, broken)

    assert not outcome["imported"] and outcome["snapshot"] is None
    assert outcome["report"]["errors"][0]["sheet"] == "columns"
    assert snapshots(roles, system) == []
    garbage = [("files", ("book.xlsx", b"PK not a workbook", "application/octet-stream"))]
    assert not upload(roles, system, garbage)["imported"]
    assert snapshots(roles, system) == []


def test_validate_reports_without_saving(roles: RoleClients):
    system = add_system(roles)

    response = roles.client("editor").post(f"{system}/import/validate", files=csvs())

    assert response.status_code == 200
    assert response.json()["errors"] == [] and response.json()["table_count"] == 1
    assert snapshots(roles, system) == []


def test_warnings_save_only_when_accepted(roles: RoleClients):
    system = add_system(roles)
    with_warning = csvs(
        constraints=(
            "schema,table,constraint,type,columns,ref_schema,ref_table,ref_columns\n"
            "sales,Customer,pk,pk,id,,,\n"
            "sales,Customer,fk,fk,id,other,thing,id\n"
        )
    )

    held = upload(roles, system, with_warning)
    assert not held["imported"] and len(held["report"]["warnings"]) == 1
    assert snapshots(roles, system) == []

    accepted = upload(roles, system, with_warning, accept_warnings=True)
    assert accepted["imported"] and accepted["snapshot"]["origin"] == "import"


def test_importing_the_same_thing_again_creates_no_snapshot(roles: RoleClients):
    system = add_system(roles)
    upload(roles, system, csvs())

    again = upload(roles, system, csvs())

    assert again["imported"] and again["unchanged"] and again["snapshot"] is None
    assert len(snapshots(roles, system)) == 1


def test_a_re_import_keeps_identities_and_is_diffed(roles: RoleClients):
    system = add_system(roles)
    upload(roles, system, csvs())
    first = latest(roles, system)
    changed = COLUMNS.replace("varchar(40)", "varchar(80)") + "sales,Customer,email,text\n"

    upload(roles, system, csvs(columns=changed))

    second = latest(roles, system)
    before, after = (
        tables_of(first)[("sales", "Customer")],
        tables_of(second)[("sales", "Customer")],
    )
    assert before["id"] == after["id"]
    response = roles.client("viewer").get(f"{system}/snapshots/{second['id']}/diff/{first['id']}")
    [table] = response.json()["tables"]
    assert table["change"] == "changed"
    assert {c["name"]: c["change"] for c in table["columns"]} == {
        "email": "added",
        "name": "changed",
    }


def test_names_match_case_insensitively_when_the_engine_folds_case(roles: RoleClients):
    system = add_system(roles)
    upload(roles, system, csvs(), engine="oracle")
    first = tables_of(latest(roles, system))[("sales", "Customer")]
    shouted = csvs(
        tables=TABLES.replace("sales,Customer", "SALES,CUSTOMER"),
        columns=COLUMNS.replace("sales,Customer", "SALES,CUSTOMER"),
    )

    upload(roles, system, shouted, engine="oracle")

    [(key, table)] = tables_of(latest(roles, system)).items()
    assert key == ("SALES", "CUSTOMER") and table["id"] == first["id"]


def test_switching_between_import_and_a_connection_keeps_identities(
    roles: RoleClients, scratch_source
):
    scratch_source["run"]("CREATE TABLE orders (id integer PRIMARY KEY, note text);")
    system = add_system(roles)
    connect(roles, system, scratch_source["body"]())
    extract(roles, system)
    live = tables_of(latest(roles, system))[("public", "orders")]

    imported = upload(
        roles,
        system,
        csvs(
            tables="schema,table,kind\npublic,orders,table\n",
            columns="schema,table,column,data_type\npublic,orders,id,integer\n"
            "public,orders,note,text\npublic,orders,extra,text\n",
        ),
    )

    assert imported["snapshot"]["origin"] == "import"
    assert tables_of(latest(roles, system))[("public", "orders")]["id"] == live["id"]


def test_an_imported_system_lists_what_does_not_work(roles: RoleClients):
    system = add_system(roles)
    before = roles.client("viewer").get(f"{system}/import").json()
    assert before["imported"] is False and before["disabled_features"] == []

    upload(roles, system, csvs())

    status = roles.client("viewer").get(f"{system}/import").json()
    assert status["imported"] and status["latest_origin"] == "import"
    assert not status["has_connection"]
    text = " ".join(status["disabled_features"]).lower()
    for feature in ("profiling", "pii", "overlap", "ai"):
        assert feature in text


def test_an_editor_requests_a_connection_and_the_owner_is_notified(roles: RoleClients):
    system = add_system(roles)
    upload(roles, system, csvs())

    response = roles.client("editor").post(f"{system}/import/connection-requests")

    assert response.status_code == 202 and response.json()["owners_notified"] >= 1
    unread = roles.client("owner").get("/api/v1/notifications").json()
    assert any(n["kind"] == "needs_owner" for n in unread["items"])


def test_a_connection_request_is_refused_once_there_is_a_connection(
    roles: RoleClients, scratch_source
):
    system = add_system(roles)
    connect(roles, system, scratch_source["body"]())

    response = roles.client("editor").post(f"{system}/import/connection-requests")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "connection_exists"
