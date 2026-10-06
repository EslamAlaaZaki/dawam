"""Reading and validating a Schema Import upload (spec §6.4, stories 47, 49).

The seam is ``read_catalog``: bytes of an upload in, a catalog and a validation report
out. No database is involved.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from dawam.modules.sources.internal.schema_import import (
    SHEETS,
    build_template,
    read_catalog,
)

TABLES = (
    "schema,table,kind,row_count,comment,definition\n"
    "sales,customer,table,10,People,\n"
    "sales,v_c,view,,,select 1\n"
)
COLUMNS = (
    "schema,table,column,data_type,ordinal,is_nullable,default,comment\n"
    "sales,customer,id,integer,1,false,,\n"
    "sales,customer,name,varchar(40),2,true,,\n"
    "sales,v_c,id,integer,,,,\n"
)
CONSTRAINTS = (
    "schema,table,constraint,type,columns,ref_schema,ref_table,ref_columns\n"
    "sales,customer,pk_customer,pk,id,,,\n"
    "sales,customer,fk_self,fk,id,sales,customer,id\n"
)
INDEXES = "schema,table,index,columns,is_unique\nsales,customer,ix_name,name,true\n"
ROUTINES = "schema,name,kind,signature,definition\nsales,f,function,(a integer),select 1\n"


def csvs(**sheets: str) -> list[tuple[str, bytes]]:
    return [(f"{name}.csv", text.encode()) for name, text in sheets.items()]


def good() -> dict[str, str]:
    return {
        "tables": TABLES,
        "columns": COLUMNS,
        "constraints": CONSTRAINTS,
        "indexes": INDEXES,
        "routines": ROUTINES,
    }


def test_csv_sheets_become_a_catalog():
    catalog, report = read_catalog(csvs(**good()))

    assert report.errors == [] and report.warnings == []
    assert catalog is not None
    assert catalog.schemas == ("sales",)
    tables = {t.name: t for t in catalog.tables}
    customer = tables["customer"]
    assert customer.kind == "table" and customer.row_estimate == 10
    assert [(c.name, c.is_pk, c.is_nullable) for c in customer.columns] == [
        ("id", True, False),
        ("name", False, True),
    ]
    assert customer.constraints[1].ref_table == "customer"
    assert customer.indexes[0].is_unique
    assert tables["v_c"].definition == "select 1"
    assert catalog.routines[0].signature == "(a integer)"
    assert (report.table_count, report.column_count, report.routine_count) == (2, 3, 1)


def test_the_filled_template_workbook_is_read_like_csvs():
    workbook = load_workbook(io.BytesIO(build_template()))
    for name, text in good().items():
        sheet = workbook[name]
        sheet.delete_rows(1, sheet.max_row)
        for line in text.splitlines():
            sheet.append(line.split(","))
    out = io.BytesIO()
    workbook.save(out)

    catalog, report = read_catalog([("filled.xlsx", out.getvalue())])

    assert report.errors == [] and report.warnings == []
    from_csv, _ = read_catalog(csvs(**good()))
    assert catalog == from_csv


def test_the_template_has_every_sheet_with_headers_and_a_readme():
    workbook = load_workbook(io.BytesIO(build_template()))

    assert workbook.sheetnames[0] == "README"
    for sheet in SHEETS:
        assert [c.value for c in workbook[sheet.name][1]] == list(sheet.column_names)
    documented = {(row[0], row[1]) for row in workbook["README"].iter_rows(values_only=True)}
    assert all((s.name, c.name) in documented for s in SHEETS for c in s.columns)


def test_an_empty_template_is_valid_but_empty():
    catalog, report = read_catalog([("template.xlsx", build_template())])

    assert report.errors == []
    assert catalog is not None and catalog.tables == ()


def test_a_missing_required_column_is_an_error_and_gives_no_catalog():
    broken = good() | {"columns": "schema,table,column\nsales,customer,id\n"}

    catalog, report = read_catalog(csvs(**broken))

    assert catalog is None
    assert [(e.sheet, e.row) for e in report.errors] == [("columns", 1)]
    assert "data_type" in report.errors[0].message


def test_a_required_sheet_that_is_missing_is_an_error():
    catalog, report = read_catalog(csvs(tables=TABLES))

    assert catalog is None
    assert [e.sheet for e in report.errors] == ["columns"]


def test_bad_values_are_reported_with_their_row():
    broken = good() | {
        "tables": TABLES + "sales,t3,matview,,,\nsales,t4,table,many,,\n",
        "columns": COLUMNS + "sales,customer,x,integer,0,maybe,,\n",
    }

    catalog, report = read_catalog(csvs(**broken))

    assert catalog is None
    problems = {(e.sheet, e.row) for e in report.errors}
    assert ("tables", 4) in problems  # kind
    assert ("tables", 5) in problems  # row_count
    assert ("columns", 5) in problems  # ordinal and is_nullable


def test_referential_consistency_is_checked():
    broken = good() | {
        "columns": COLUMNS + "sales,ghost,id,integer,,,,\n",
        "constraints": CONSTRAINTS
        + "sales,customer,pk_bad,unique,nope,,,\n"
        + "sales,customer,fk_bad,fk,id,sales,customer,missing\n",
        "indexes": INDEXES + "sales,ghost,ix,id,false\n",
    }

    _, report = read_catalog(csvs(**broken))

    messages = " ".join(e.message for e in report.errors)
    assert "sales.ghost" in messages
    assert "nope" in messages
    assert "missing" in messages
    assert {e.sheet for e in report.errors} == {"columns", "constraints", "indexes"}


def test_a_foreign_key_to_a_table_outside_the_import_is_a_warning():
    with_warning = good() | {
        "constraints": CONSTRAINTS + "sales,customer,fk_out,fk,id,other,thing,id\n"
    }

    catalog, report = read_catalog(csvs(**with_warning))

    assert report.errors == []
    assert catalog is not None
    assert len(report.warnings) == 1 and "other.thing" in report.warnings[0].message
    assert not report.clean


def test_duplicates_are_errors():
    broken = good() | {"tables": TABLES + "sales,customer,table,,,\n"}

    _, report = read_catalog(csvs(**broken))

    assert [(e.sheet, e.row) for e in report.errors] == [("tables", 4)]


def test_unreadable_files_are_reported_not_raised():
    for name, content in [
        ("broken.xlsx", b"PK this is not a workbook"),
        ("notes.txt", b"hello"),
        ("whatever.csv", b"x"),
        ("tables.csv", b"\xff\xfe\x00bad"),
    ]:
        catalog, report = read_catalog([(name, content)])
        assert catalog is None
        assert report.errors, name


def test_the_same_sheet_twice_is_an_error():
    _, report = read_catalog([*csvs(**good()), ("tables.csv", TABLES.encode())])

    assert any("twice" in e.message for e in report.errors)
