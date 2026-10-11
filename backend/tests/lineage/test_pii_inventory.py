"""PII inventory export (spec §6.12, story 140): XLSX and Markdown, PII-influenced tables
listed apart from the PII columns, download or save to the file area."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import sqlalchemy as sa
from openpyxl import load_workbook

from tests.lineage import test_lineage
from tests.lineage.test_lineage import dw, ok
from tests.roles import RoleClients
from tests.warehouse.test_pii_hashing import hash_column, set_note

graph = test_lineage.graph  # the lineage suite's Workspace (a fixture)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_url(roles: RoleClients) -> str:
    return f"/api/v1/workspaces/{roles.workspace_id}/pii/inventory/export"


def confirm(roles: RoleClients, column_id: str) -> None:
    """A confirmed finding decided by the editor."""
    with roles.app.state.engine.begin() as connection:
        connection.execute(
            sa.text(
                "insert into pii_findings (id, src_column_id, rule, category, confidence,"
                " evidence, status, detected_at, decided_by, decided_at) values (:id, :c,"
                " 'national_id', 'direct_identifier', 0.9, 'name', 'confirmed', :now, :by, :now)"
            ),
            {
                "id": uuid.uuid4(),
                "c": column_id,
                "now": datetime(2026, 10, 1, tzinfo=UTC),
                "by": roles.user("editor").id,
            },
        )


def book(data: bytes) -> dict[str, list[dict]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    result = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        result[sheet.title] = [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]
    return result


def test_a_member_downloads_the_inventory_as_xlsx(graph, roles):
    confirm(roles, graph["src_amount"])
    stg = graph["stg"]["name"]
    hash_column(roles, graph["stg_amount"])
    set_note(roles, graph["mart"], graph["amount_m"], "Aggregate before sharing.")

    response = roles.client("viewer").get(export_url(roles))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    assert "pii-inventory.xlsx" in response.headers["content-disposition"]
    sheets = book(response.content)
    assert list(sheets) == ["Source columns", "DW columns", "Influenced tables"]
    [source] = sheets["Source columns"]
    assert (source["source_system"], source["database_schema"], source["table"]) == (
        "crm",
        "dbo",
        "orders",
    )
    assert source["column"] == "amount"
    assert source["category"] == "direct_identifier"
    assert source["reviewer"] == roles.user("editor").display_name
    assert source["reviewed_at"] == "2026-10-01"
    columns = {(r["table"], r["column"]): r for r in sheets["DW columns"]}
    assert set(columns) == {
        (stg, "amount"),
        ("fact_orders", "amount_c"),
        ("mart_orders", "amount_m"),
    }
    mart = columns[("mart_orders", "amount_m")]
    assert mart["layer"] == "mart"
    assert mart["source_columns"] == "crm.dbo.orders.amount"
    assert mart["category"] == "direct_identifier"
    assert mart["hashing"] == "pseudonymised"
    assert mart["handling_note"] == "Aggregate before sharing."
    assert mart["reviewer"] == roles.user("editor").display_name


def test_influenced_tables_are_listed_apart_from_pii_columns(graph, roles):
    confirm(roles, graph["src_status"])

    sheets = book(roles.client("viewer").get(export_url(roles)).content)

    assert [r["column"] for r in sheets["DW columns"]] == ["status"]
    assert [(r["layer"], r["table"]) for r in sheets["Influenced tables"]] == [
        ("core", "fact_orders")
    ]
    assert all(r["table"] != "fact_orders" for r in sheets["DW columns"])


def test_a_column_that_is_not_hashed_says_so(graph, roles):
    confirm(roles, graph["src_amount"])

    sheets = book(roles.client("viewer").get(export_url(roles)).content)

    assert {r["hashing"] for r in sheets["DW columns"]} == {"not hashed"}


def test_a_member_downloads_the_inventory_as_markdown(graph, roles):
    confirm(roles, graph["src_status"])

    response = roles.client("viewer").get(export_url(roles), params={"format": "md"})

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/markdown")
    assert "pii-inventory.md" in response.headers["content-disposition"]
    text = response.text
    assert text.startswith("# PII inventory")
    assert "## Source columns" in text
    assert "| crm | dbo | orders | status | direct_identifier |" in text
    assert text.index("## DW columns") < text.index("## Influenced tables")
    influenced = text[text.index("## Influenced tables") :]
    assert "| core | fact_orders |" in influenced


def test_an_empty_inventory_still_exports(graph, roles):
    response = roles.client("viewer").get(export_url(roles), params={"format": "md"})

    assert response.status_code == 200, response.text
    assert response.text.count("None.") == 3
    sheets = book(roles.client("viewer").get(export_url(roles)).content)
    assert all(rows == [] for rows in sheets.values())


def test_unknown_formats_are_refused(graph, roles):
    response = roles.client("viewer").get(export_url(roles), params={"format": "csv"})

    assert response.status_code == 422


def test_markdown_cells_cannot_break_the_table(graph, roles):
    confirm(roles, graph["src_amount"])
    set_note(roles, graph["mart"], graph["amount_m"], "a | b")

    text = roles.client("viewer").get(export_url(roles), params={"format": "md"}).text

    assert "a \\| b" in text


def test_the_inventory_is_saved_to_the_file_area(graph, roles):
    confirm(roles, graph["src_amount"])
    save = f"{dw(roles)}/files/pii-inventory"

    saved = roles.client("editor").post(save, params={"format": "md"})

    assert saved.status_code == 201, saved.text
    assert saved.json()["name"] == "pii-inventory.md"
    again = roles.client("editor").post(save, params={"format": "md"})
    assert again.status_code == 201
    listing = ok(roles.client("viewer").get(f"{dw(roles)}/files"))["items"]
    assert [f["name"] for f in listing].count("pii-inventory.md") == 1
    body = roles.client("viewer").get(
        f"/api/v1/workspaces/{roles.workspace_id}/files/{saved.json()['id']}/download"
    )
    assert "# PII inventory" in body.text


def test_a_viewer_cannot_save_the_inventory(graph, roles):
    response = roles.client("viewer").post(f"{dw(roles)}/files/pii-inventory")

    assert response.status_code == 403
