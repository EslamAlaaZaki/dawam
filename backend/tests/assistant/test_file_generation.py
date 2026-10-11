"""Assistant file generation and updates (spec stories 145, 146): ``generate_file`` saves
exporter output (DDL, XLSX) or free text in a file area; ``update_file`` proposes a new text as
a Change Set that changes the file only once accepted. The model is the scripted fake."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from tests.assistant import test_chat
from tests.assistant.test_source_tools import payload, use
from tests.kpis import test_kpi_links
from tests.roles import RoleClients
from tests.warehouse.test_ddl import modelled

model = test_chat.model  # the assistant's model (a fixture)
warehouse = test_kpi_links.model  # a Data Warehouse with Core and Mart tables (a fixture)
__all__ = ["model", "modelled", "warehouse"]


def base(roles: RoleClients) -> str:
    return f"/api/v1/workspaces/{roles.workspace_id}"


def area_files(roles: RoleClients) -> list[dict[str, Any]]:
    response = roles.client("owner").get(f"{base(roles)}/data-warehouse/files")
    assert response.status_code == 200, response.text
    return response.json()["items"]


def text_of(roles: RoleClients, file: dict[str, Any]) -> str:
    response = roles.client("owner").get(f"{base(roles)}/files/{file['id']}/content")
    assert response.status_code == 200, response.text
    return response.json()["content"]


def a_file(roles: RoleClients, text: str = "# Notes\n\nold\n") -> dict[str, Any]:
    response = roles.client("owner").post(
        f"{base(roles)}/data-warehouse/files",
        files={"file": ("notes.md", text.encode(), "text/markdown")},
    )
    assert response.status_code in (200, 201), response.text
    return response.json()


def change_set_ids(roles: RoleClients) -> list[str]:
    listing = roles.client("editor").get(f"{base(roles)}/change-sets").json()
    return [c["id"] for c in listing["items"]]


def decide(roles: RoleClients, change_set: str, verb: str):
    return roles.client("editor").post(f"{base(roles)}/change-sets/{change_set}/{verb}", json={})


# -- generate_file ---------------------------------------------------------------------------


def test_ddl_is_saved_to_the_warehouse_area_through_the_exporter(roles, model, modelled, fake_llm):
    status, _ = use(roles, fake_llm, "generate_file", "editor", kind="ddl")

    assert status == "ok"
    [saved] = area_files(roles)
    assert (saved["name"], saved["kind"]) == ("ddl-data-warehouse.sql", "generated")
    exported = roles.client("viewer").get(f"{base(roles)}/data-warehouse/ddl").text
    assert text_of(roles, saved) == exported


def test_the_mapping_sheet_is_saved_as_the_exporters_xlsx(roles, model, warehouse, fake_llm):
    status, _ = use(roles, fake_llm, "generate_file", "editor", kind="mapping_sheet")

    assert status == "ok"
    saved = next(f for f in area_files(roles) if f["name"].endswith(".xlsx"))
    data = roles.client("viewer").get(f"{base(roles)}/files/{saved['id']}/download").content
    assert "Mapping" in load_workbook(io.BytesIO(data)).sheetnames


def test_free_text_goes_to_the_warehouse_area_without_a_system(roles, model, warehouse, fake_llm):
    status, _ = use(
        roles, fake_llm, "generate_file", "editor", name="kpi-docs.md", content="# KPIs\n"
    )

    assert status == "ok"
    [saved] = area_files(roles)
    assert text_of(roles, saved) == "# KPIs\n"


def test_regenerating_an_export_asks_before_overwriting(roles, model, modelled, fake_llm):
    use(roles, fake_llm, "generate_file", "editor", kind="ddl")

    _, seen = use(roles, fake_llm, "generate_file", "editor", kind="ddl")

    assert payload(seen)["status"] in ("unchanged", "needs_confirmation")


def test_a_viewer_cannot_generate_a_file(roles, model, modelled, fake_llm):
    status, _ = use(roles, fake_llm, "generate_file", "viewer", kind="ddl")

    assert status == "refused"
    assert area_files(roles) == []


# -- update_file -----------------------------------------------------------------------------


def test_an_update_is_a_change_set_with_a_diff_and_the_file_is_unchanged(
    roles, model, warehouse, fake_llm
):
    file = a_file(roles)

    status, seen = use(
        roles, fake_llm, "update_file", "editor", file_id=file["id"], content="# Notes\n\nnew\n"
    )

    assert status == "ok"
    body = payload(seen)
    assert body["status"] == "proposed" and "+new" in body["diff"] and "-old" in body["diff"]
    assert text_of(roles, file) == "# Notes\n\nold\n"
    [proposed] = roles.client("editor").get(f"{base(roles)}/change-sets").json()["items"]
    assert proposed["origin"] == "ai" and proposed["status"] == "pending"


def test_accepting_saves_the_new_text(roles, model, warehouse, fake_llm):
    file = a_file(roles)
    use(roles, fake_llm, "update_file", "editor", file_id=file["id"], content="# Notes\n\nnew\n")
    [change_set] = change_set_ids(roles)

    accepted = decide(roles, change_set, "accept")

    assert accepted.status_code == 200, accepted.text
    assert text_of(roles, file) == "# Notes\n\nnew\n"


def test_rejecting_leaves_the_file_as_it_was(roles, model, warehouse, fake_llm):
    file = a_file(roles)
    use(roles, fake_llm, "update_file", "editor", file_id=file["id"], content="# Notes\n\nnew\n")
    [change_set] = change_set_ids(roles)

    rejected = decide(roles, change_set, "reject")

    assert rejected.status_code == 200, rejected.text
    assert text_of(roles, file) == "# Notes\n\nold\n"


def test_an_edit_made_after_the_proposal_makes_it_stale(roles, model, warehouse, fake_llm):
    file = a_file(roles)
    use(roles, fake_llm, "update_file", "editor", file_id=file["id"], content="# Notes\n\nnew\n")
    [change_set] = change_set_ids(roles)
    edited = roles.client("editor").put(
        f"{base(roles)}/files/{file['id']}/content", json={"content": "# Notes\n\nmine\n"}
    )
    assert edited.status_code == 200, edited.text

    accepted = decide(roles, change_set, "accept")

    assert accepted.status_code == 200, accepted.text
    assert [s["reason"] for s in accepted.json()["skipped"]] == ["stale"]
    assert text_of(roles, file) == "# Notes\n\nmine\n"


def test_a_viewer_cannot_update_a_file(roles, model, warehouse, fake_llm):
    file = a_file(roles)

    status, _ = use(
        roles, fake_llm, "update_file", "viewer", file_id=file["id"], content="# Notes\n\nnew\n"
    )

    assert status == "refused"
    assert change_set_ids(roles) == []
    assert text_of(roles, file) == "# Notes\n\nold\n"


def test_a_binary_file_cannot_be_updated(roles, model, warehouse, fake_llm):
    use(roles, fake_llm, "generate_file", "editor", kind="mapping_sheet")
    sheet = next(f for f in area_files(roles) if f["name"].endswith(".xlsx"))

    status, _ = use(roles, fake_llm, "update_file", "editor", file_id=sheet["id"], content="x")

    assert status == "error"
