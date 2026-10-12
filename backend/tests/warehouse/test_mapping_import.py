"""Mapping sheet import (spec §6.14, story 109): a validation report, then a Change Set."""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from dawam.modules.warehouse.mapping_export import BRANCH_HEADER, MAPPING_HEADER
from tests.roles import RoleClients
from tests.warehouse.test_column_mappings import base, get_mapping
from tests.warehouse.test_mapping_branches import model, two_branches, warehouse

__all__ = ["model", "warehouse"]


def sheet_bytes(mapping: list[dict], branches: list[dict] | None = None) -> bytes:
    """A mapping sheet from rows given as ``{header: value}``."""
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Mapping"
    first.append(MAPPING_HEADER)
    for row in mapping:
        first.append([row.get(h) for h in MAPPING_HEADER])
    if branches is not None:
        second = workbook.create_sheet("Branches")
        second.append(BRANCH_HEADER)
        for row in branches:
            second.append([row.get(h) for h in BRANCH_HEADER])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def exported(roles: RoleClients) -> tuple[list[dict], list[dict]]:
    response = roles.client("viewer").get(f"{base(roles)}/mapping-sheet")
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    sheets = {}
    for ws in workbook.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        sheets[ws.title] = [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]
    return sheets["Mapping"], sheets["Branches"]


def upload(roles: RoleClients, data: bytes, *, dry_run: bool = True, role: str = "editor"):
    return roles.client(role).post(
        f"{base(roles)}/mapping-sheet/import",
        params={"dry_run": dry_run},
        files={"file": ("mapping-sheet.xlsx", data, "application/octet-stream")},
    )


def change_set(roles: RoleClients, change_set_id: str) -> dict:
    response = roles.client("viewer").get(
        f"/api/v1/workspaces/{roles.workspace_id}/change-sets/{change_set_id}"
    )
    assert response.status_code == 200, response.text
    return response.json()


def tables(roles: RoleClients) -> list[str]:
    response = roles.client("viewer").get(f"{base(roles)}/tables")
    assert response.status_code == 200, response.text
    return sorted(t["name"] for t in response.json()["items"])


def test_exporting_and_importing_again_proposes_no_changes(warehouse, model):
    two_branches(warehouse, model)
    mapping, branches = exported(warehouse)

    response = upload(warehouse, sheet_bytes(mapping, branches), dry_run=False)

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["report"]["errors"] == []
    assert body["report"]["changes"] == 0
    assert body["report"]["unchanged"] == len(mapping) + len(branches)
    assert body["change_set_id"] is None


def test_rows_for_missing_tables_columns_or_branches_fail_and_create_nothing(warehouse, model):
    two_branches(warehouse, model)
    mapping, _ = exported(warehouse)
    [name_row] = [r for r in mapping if r["branch"] == "CRM" and r["target_column"] == "name"]
    ghost_table = {**name_row, "target_table": "ghost", "sql_expression": "a.b"}
    ghost_column = {**name_row, "target_column": "ghost", "sql_expression": "crm_customer.x"}
    ghost_branch = {**name_row, "branch": "Ghost", "sql_expression": "crm_customer.x"}
    before = tables(warehouse)
    mart_before = get_mapping(warehouse, model["mart"]["id"])

    response = upload(
        warehouse, sheet_bytes([ghost_table, ghost_column, ghost_branch]), dry_run=False
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert [e["row"] for e in body["report"]["errors"]] == [2, 3, 4]
    assert all(e["sheet"] == "Mapping" for e in body["report"]["errors"])
    assert body["change_set_id"] is None
    assert tables(warehouse) == before
    assert get_mapping(warehouse, model["mart"]["id"]) == mart_before


def test_a_missing_branch_on_the_branches_sheet_fails_too(warehouse, model):
    two_branches(warehouse, model)
    _, branches = exported(warehouse)

    response = upload(
        warehouse, sheet_bytes([], [{**branches[0], "branch": "Ghost"}]), dry_run=False
    )

    [error] = response.json()["report"]["errors"]
    assert (error["sheet"], error["row"]) == ("Branches", 2)
    assert response.json()["change_set_id"] is None


def test_staging_rows_are_ignored(warehouse, model):
    row = {"layer": "staging", "target_table": "stg_ghost", "target_column": "x", "branch": ""}

    response = upload(warehouse, sheet_bytes([row]), dry_run=False)

    report = response.json()["report"]
    assert report["errors"] == [] and report["ignored_staging_rows"] == 1
    assert report["changes"] == 0


def test_a_changed_mapping_becomes_a_pending_import_change_set_and_nothing_else(warehouse, model):
    two_branches(warehouse, model)
    mapping, _ = exported(warehouse)
    [row] = [r for r in mapping if r["branch"] == "CRM" and r["target_column"] == "name"]
    changed = {
        **row,
        "mapping_type": "derived",
        "sql_expression": "UPPER(crm_customer.full_name)",
        "transformation_rule": "Upper-case the name",
    }
    approvals = warehouse.client("viewer").get(f"{base(warehouse)}/approvals").json()
    before = get_mapping(warehouse, model["mart"]["id"])

    dry = upload(warehouse, sheet_bytes([changed]))
    response = upload(warehouse, sheet_bytes([changed]), dry_run=False)

    assert dry.json()["report"]["changes"] == 1 and dry.json()["change_set_id"] is None
    body = response.json()
    assert body["report"]["errors"] == [] and body["report"]["changes"] == 1
    detail = change_set(warehouse, body["change_set_id"])
    assert detail["change_set"]["origin"] == "import"
    assert detail["change_set"]["status"] == "pending"
    [item] = detail["items"]
    assert (item["object_type"], item["operation"]) == ("column_mapping", "update")
    assert item["payload"]["sql_expression"] == "UPPER(crm_customer.full_name)"
    assert get_mapping(warehouse, model["mart"]["id"]) == before
    assert warehouse.client("viewer").get(f"{base(warehouse)}/approvals").json() == approvals

    accepted = warehouse.client("editor").post(
        f"/api/v1/workspaces/{warehouse.workspace_id}/change-sets/{body['change_set_id']}/accept",
        json={},
    )
    assert accepted.status_code == 200, accepted.text
    crm = get_mapping(warehouse, model["mart"]["id"])["branches"][0]
    [name] = [c for c in crm["columns"] if c["column_name"] == "name"]
    assert name["mapping_type"] == "derived"
    assert name["sql_expression"] == "UPPER(crm_customer.full_name)"
    assert name["rule_text"] == "Upper-case the name"


def test_an_unmapped_column_is_mapped_by_a_created_item(warehouse, model):
    row = {
        "layer": "mart",
        "target_table": "dim_customer",
        "branch": "",
        "target_column": "name",
        "mapping_type": "direct",
        "sql_expression": "crm_customer.full_name",
    }

    response = upload(warehouse, sheet_bytes([row]), dry_run=False)

    body = response.json()
    assert body["report"]["errors"] == [], body
    [item] = change_set(warehouse, body["change_set_id"])["items"]
    assert (item["object_type"], item["operation"]) == ("column_mapping", "create")
    accepted = warehouse.client("editor").post(
        f"/api/v1/workspaces/{warehouse.workspace_id}/change-sets/{body['change_set_id']}/accept",
        json={},
    )
    assert accepted.status_code == 200, accepted.text
    mapped = {c["column_name"]: c for c in get_mapping(warehouse, model["mart"]["id"])["columns"]}
    assert mapped["name"]["mapping_type"] == "direct"
    assert mapped["name"]["sql_expression"] == "crm_customer.full_name"


def test_sql_that_names_a_missing_column_fails_validation(warehouse, model):
    row = {
        "layer": "mart",
        "target_table": "dim_customer",
        "target_column": "name",
        "mapping_type": "direct",
        "sql_expression": "crm_customer.no_such_column",
    }

    response = upload(warehouse, sheet_bytes([row]), dry_run=False)

    body = response.json()
    assert [e["row"] for e in body["report"]["errors"]] == [2]
    assert body["change_set_id"] is None


def test_a_mapping_the_tool_makes_cannot_be_replaced_by_an_import(warehouse, model):
    two_branches(warehouse, model)
    mapping, _ = exported(warehouse)
    [system] = [
        r for r in mapping if r["mapping_type"] == "system" and r["target_table"] == "dim_customer"
    ]

    response = upload(
        warehouse,
        sheet_bytes([{**system, "mapping_type": "derived", "sql_expression": "1"}]),
        dry_run=False,
    )

    assert len(response.json()["report"]["errors"]) == 1
    assert response.json()["change_set_id"] is None


def test_a_file_that_is_not_a_mapping_sheet_is_refused(warehouse, model):
    garbage = upload(warehouse, b"not a workbook")
    empty = Workbook()
    buffer = io.BytesIO()
    empty.save(buffer)
    no_sheet = upload(warehouse, buffer.getvalue())

    assert garbage.status_code == 422 and garbage.json()["error"]["code"] == "invalid_sheet"
    assert no_sheet.status_code == 422 and no_sheet.json()["error"]["code"] == "invalid_sheet"


def test_a_viewer_cannot_import(warehouse, model):
    response = upload(warehouse, sheet_bytes([]), role="viewer")

    assert response.status_code == 403
