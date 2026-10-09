"""Mapping sheet export (spec §6.14, story 108): XLSX and CSV in the mapping-sheet layout."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from tests.roles import RoleClients
from tests.warehouse.test_column_mappings import base, mapping_path
from tests.warehouse.test_mapping_branches import model, put_branch_column, two_branches, warehouse

__all__ = ["model", "warehouse"]

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def sheet_url(roles: RoleClients) -> str:
    return f"{base(roles)}/mapping-sheet"


def book(data: bytes) -> dict[str, list[dict]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    result = {}
    for ws in workbook.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        result[ws.title] = [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]
    return result


def merged(roles, model) -> None:
    two_branches(roles, model)
    response = roles.client("editor").patch(
        mapping_path(roles, model["mart"]["id"]),
        json={
            "version": 0,
            "integration_rule": "Merge on email, CRM wins",
            "match_keys": ["email"],
            "notes": "Customer golden record",
        },
    )
    assert response.status_code == 200, response.text


def test_a_member_downloads_the_mapping_sheet_with_a_row_per_branch_column(warehouse, model):
    merged(warehouse, model)

    response = warehouse.client("viewer").get(sheet_url(warehouse))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    assert "mapping-sheet-data-warehouse.xlsx" in response.headers["content-disposition"]
    sheets = book(response.content)
    assert list(sheets) == ["Mapping", "Branches"]
    rows = [r for r in sheets["Mapping"] if r["target_table"] == "dim_customer"]
    assert len(rows) == 7  # 3 columns x 2 branches, plus the surrogate key
    [cb_name] = [r for r in rows if r["branch"] == "Core Banking" and r["target_column"] == "name"]
    assert cb_name["layer"] == "mart"
    assert cb_name["input_layer"] == "core"
    assert cb_name["input_table"] == "cb_customer"
    assert cb_name["input_column(s)"] == "cb_customer.full_name"
    assert cb_name["sql_expression"] == "cb_customer.full_name"
    assert cb_name["mapping_type"] == "direct"
    assert cb_name["target_type"]
    assert cb_name["notes"] == "Customer golden record"
    [crm_email] = [r for r in rows if r["branch"] == "CRM" and r["target_column"] == "email"]
    assert crm_email["mapping_type"] == "not_in_branch"
    assert crm_email["input_column(s)"] in (None, "")


def test_several_inputs_are_one_cell_of_table_dot_column_joined_by_semicolons(warehouse, model):
    crm, _ = two_branches(warehouse, model)
    response = put_branch_column(
        warehouse,
        model,
        crm,
        "name",
        version=1,
        mapping_type="derived",
        sql_expression="crm_customer.full_name || crm_customer.customer_id",
    )
    assert response.status_code == 200, response.text

    rows = book(warehouse.client("viewer").get(sheet_url(warehouse)).content)["Mapping"]

    [row] = [r for r in rows if r["branch"] == "CRM" and r["target_column"] == "name"]
    assert row["input_column(s)"] == "crm_customer.customer_id;crm_customer.full_name"
    assert row["input_table"] == "crm_customer"


def test_the_second_sheet_holds_each_branchs_query_parts_and_the_integration_rule(warehouse, model):
    merged(warehouse, model)

    sheets = book(warehouse.client("viewer").get(sheet_url(warehouse)).content)

    branches = sheets["Branches"]
    assert [(b["target_table"], b["branch"], b["driving_input"]) for b in branches] == [
        ("dim_customer", "CRM", "crm_customer"),
        ("dim_customer", "Core Banking", "cb_customer"),
    ]
    assert all(b["integration_rule"] == "Merge on email, CRM wins" for b in branches)
    assert all(b["match_keys"] == "email" for b in branches)
    assert {"joins", "filters", "group_by", "having"} <= set(branches[0])


def test_csv_holds_the_mapping_columns_and_the_layer_filter_applies(warehouse, model):
    merged(warehouse, model)

    response = warehouse.client("viewer").get(
        sheet_url(warehouse), params={"format": "csv", "layer": "mart"}
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    assert "mapping-sheet-mart.csv" in response.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert list(rows[0]) == [
        "layer",
        "target_table",
        "branch",
        "target_column",
        "target_type",
        "input_layer",
        "input_system",
        "input_database_schema",
        "input_table",
        "input_column(s)",
        "transformation_rule",
        "sql_expression",
        "mapping_type",
        "lookup_dimension",
        "notes",
    ]
    assert {r["layer"] for r in rows} == {"mart"}


def test_a_formula_like_rule_stays_text_in_the_workbook(warehouse, model):
    crm, _ = two_branches(warehouse, model)
    response = put_branch_column(
        warehouse,
        model,
        crm,
        "name",
        version=1,
        mapping_type="direct",
        sql_expression="crm_customer.full_name",
        rule_text="=1+1",
    )
    assert response.status_code == 200, response.text

    rows = book(warehouse.client("viewer").get(sheet_url(warehouse)).content)["Mapping"]

    [row] = [r for r in rows if r["branch"] == "CRM" and r["target_column"] == "name"]
    assert row["transformation_rule"] == "=1+1"


def test_csv_cells_that_look_like_formulas_are_prefixed_with_a_quote(warehouse, model):
    crm, _ = two_branches(warehouse, model)
    response = put_branch_column(
        warehouse,
        model,
        crm,
        "name",
        version=1,
        mapping_type="direct",
        sql_expression="crm_customer.full_name",
        rule_text='=HYPERLINK("http://evil")',
    )
    assert response.status_code == 200, response.text

    csv_response = warehouse.client("viewer").get(sheet_url(warehouse), params={"format": "csv"})

    rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    [row] = [r for r in rows if r["branch"] == "CRM" and r["target_column"] == "name"]
    assert row["transformation_rule"] == '\'=HYPERLINK("http://evil")'
    from dawam.modules.warehouse.mapping_export import _csv_safe

    for text in ("=1", "+1", "-1", "@x", "\tx", "\rx"):
        assert _csv_safe(text) == "'" + text
    assert _csv_safe("plain") == "plain"


def test_an_editor_saves_the_sheet_to_the_file_area(warehouse, model):
    merged(warehouse, model)

    saved = warehouse.client("editor").post(
        f"{base(warehouse)}/files/mapping-sheet", params={"format": "csv"}
    )

    assert saved.status_code == 201, saved.text
    assert saved.json()["name"] == "mapping-sheet-data-warehouse.csv"
    listed = warehouse.client("viewer").get(f"{base(warehouse)}/files").json()["items"]
    assert "mapping-sheet-data-warehouse.csv" in [f["name"] for f in listed]


def test_the_sheet_needs_a_set_up_warehouse_and_a_known_layer(roles: RoleClients, warehouse):
    assert (
        warehouse.client("viewer")
        .get(sheet_url(warehouse), params={"layer": "staging"})
        .status_code
        == 422
    )


def test_an_unset_warehouse_has_no_sheet(roles: RoleClients):
    assert roles.client("viewer").get(sheet_url(roles)).status_code == 404
