"""Score report export (spec §6.11, story 124): Markdown and XLSX, downloaded or saved to the
file area."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

from openpyxl import load_workbook

from dawam.modules.warehouse.score_export import ScoreExportService
from dawam.modules.warehouse.score_service import FailedCheck, LayerView, Score, TableView
from tests.roles import RoleClients
from tests.warehouse.test_column_mappings import base, column
from tests.warehouse.test_score_api import fact, warehouse

__all__ = ["warehouse"]

_ID = uuid.uuid4()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def report_url(roles: RoleClients) -> str:
    return f"{base(roles)}/score/report"


def book(data: bytes) -> dict[str, list[dict]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    result = {}
    for ws in workbook.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        result[ws.title] = [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]
    return result


def failing_design(roles: RoleClients) -> dict:
    sales = fact(roles)
    column(roles, sales["id"], "amount", "decimal", role="measure")
    return sales


def test_a_member_downloads_the_report_as_markdown(warehouse: RoleClients):
    failing_design(warehouse)

    response = warehouse.client("viewer").get(report_url(warehouse))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/markdown")
    assert 'filename="score-report.md"' in response.headers["content-disposition"]
    text = response.text
    assert "# Data Warehouse score report" in text
    for heading in ("## Overall", "## Layers", "## Tables", "## Failed checks"):
        assert heading in text
    assert "| core |" in text and "| mart |" in text and "| staging |" in text
    assert "fact_sales" in text
    assert "Fix:" in text and "ERROR:" in text


def test_the_xlsx_report_has_overall_layer_table_and_failed_check_sheets(warehouse: RoleClients):
    failing_design(warehouse)

    response = warehouse.client("viewer").get(report_url(warehouse), params={"format": "xlsx"})

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    assert 'filename="score-report.xlsx"' in response.headers["content-disposition"]
    sheets = book(response.content)
    assert list(sheets) == ["Summary", "Layers", "Tables", "Failed checks"]
    summary = {r["measure"]: r["value"] for r in sheets["Summary"]}
    assert summary["grade"] in ("C", "D", "F") and summary["checks_failed"] >= 1
    assert [r["layer"] for r in sheets["Layers"]] == ["staging", "core", "mart"]
    assert [r["table"] for r in sheets["Tables"]] == ["fact_sales"]
    [missing] = [r for r in sheets["Failed checks"] if r["check"] == "measure_has_additivity"]
    assert missing["object"] == "fact_sales.amount" and missing["fix_hint"] and missing["severity"]


def test_an_empty_model_reports_not_scored(warehouse: RoleClients):
    response = warehouse.client("viewer").get(report_url(warehouse))

    assert response.status_code == 200, response.text
    assert "not scored" in response.text
    assert "No failed checks." in response.text


def test_an_unknown_format_is_refused(warehouse: RoleClients):
    response = warehouse.client("viewer").get(report_url(warehouse), params={"format": "pdf"})

    assert response.status_code == 422


def test_the_report_needs_a_data_warehouse(roles: RoleClients):
    assert roles.client("viewer").get(report_url(roles)).status_code == 404


def test_saving_puts_the_report_in_the_file_area_and_overwrites_it(warehouse: RoleClients):
    failing_design(warehouse)
    save = f"{base(warehouse)}/files/score-report"

    first = warehouse.client("editor").post(save, params={"format": "xlsx"})
    second = warehouse.client("editor").post(save, params={"format": "xlsx"})
    markdown = warehouse.client("editor").post(save)

    assert first.status_code == 201, first.text
    assert first.json()["name"] == "score-report.xlsx"
    assert second.json()["id"] == first.json()["id"]
    assert markdown.json()["name"] == "score-report.md"


def test_cells_that_look_like_formulas_stay_text_and_pipes_are_escaped():
    now = datetime(2026, 1, 1, tzinfo=UTC)
    check = FailedCheck(
        check_code="x",
        title="a|b",
        category="c",
        severity="error",
        layer="core",
        object_type="table",
        object_id=_ID,
        table_id=_ID,
        object_name="=SUM(A1)",
        message="line one\nline two",
        fix_hint="+fix",
        link="?table=1",
    )
    score = Score(
        score=70.0,
        grade="C",
        capped=True,
        calculated_at=now,
        layers=[LayerView("core", 70.0, "C", 1)],
        tables=[TableView(_ID, "core", "=SUM(A1)", 70.0)],
        failed_checks=[check],
        checks_passed=1,
        checks_failed=1,
    )

    markdown = ScoreExportService._markdown(score)
    sheets = book(ScoreExportService._xlsx(score))

    assert "a\\|b" in markdown and "line one line two" in markdown
    assert "capped at C" in markdown
    assert sheets["Tables"][0]["table"] == "=SUM(A1)"
    assert sheets["Failed checks"][0]["fix_hint"] == "+fix"
