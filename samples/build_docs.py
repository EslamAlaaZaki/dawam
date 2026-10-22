# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx==1.1.2",
#     "openpyxl==3.1.5",
#     "fpdf2==2.8.1",
# ]
# ///
"""Render a sample scenario's documents from Markdown: ``uv run samples/build_docs.py
samples/<scenario>``.

Each ``<scenario>/docs/*.md`` whose first line is ``<!-- build: docx -->`` (or ``xlsx``,
``pdf``) is rendered next to itself with the same name and that extension; any other
Markdown file (a glossary, say) ships as it is. The Markdown is a small subset:

- ``#``, ``##`` and ``###`` headings, paragraphs, ``-`` bullets, ``1.`` numbered items
  and pipe tables, with ``**bold**`` and ```code``` inside text;
- for ``xlsx``, every ``##`` heading starts a sheet of that name, filled from the first
  pipe table under it (everything else is ignored).

The output is deterministic: document dates are fixed and the DOCX/XLSX zip entries carry
a fixed timestamp, so rebuilding unchanged sources gives byte-identical files.
"""

from __future__ import annotations

import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt
from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

FIXED_DATE = datetime(2025, 1, 1, tzinfo=timezone.utc)
"""The creation and modification date written into every document."""
AUTHOR = "DAWAM samples"
_MARKER = re.compile(r"<!--\s*build:\s*(docx|xlsx|pdf)\s*-->")


# -- Markdown --------------------------------------------------------------------------


@dataclass
class Block:
    kind: str
    """``heading``, ``paragraph``, ``bullet``, ``number`` or ``table``."""
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)
    """For a table: the header row first, without the ``---`` separator row."""


def parse(text: str) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph)))
            paragraph.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("<!--"):
            flush()
        elif heading := re.match(r"(#{1,3})\s+(.*)", line):
            flush()
            blocks.append(Block("heading", heading[2], level=len(heading[1])))
        elif line.startswith("|"):
            flush()
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            if not blocks or blocks[-1].kind != "table":
                blocks.append(Block("table"))
            blocks[-1].rows.append(cells)
        elif bullet := re.match(r"-\s+(.*)", line):
            flush()
            blocks.append(Block("bullet", bullet[1]))
        elif number := re.match(r"(\d+)\.\s+(.*)", line):
            flush()
            blocks.append(Block("number", number[2], level=int(number[1])))
        elif raw.startswith("  ") and blocks and blocks[-1].kind in ("bullet", "number"):
            blocks[-1].text += " " + line  # a list item's continuation line
        else:
            paragraph.append(line)
    flush()
    return blocks


_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def spans(text: str) -> list[tuple[str, str]]:
    """Text as (style, text) pieces; style is ``""``, ``bold`` or ``code``."""
    pieces = []
    for part in _INLINE.split(text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            pieces.append(("bold", part[2:-2]))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            pieces.append(("code", part[1:-1]))
        elif part:
            pieces.append(("", part))
    return pieces


def plain(text: str) -> str:
    return "".join(piece for _, piece in spans(text))


# -- DOCX ------------------------------------------------------------------------------


def _runs(paragraph, text: str) -> None:
    for style, piece in spans(text):
        run = paragraph.add_run(piece)
        run.bold = style == "bold" or None
        if style == "code":
            run.font.name = "Consolas"


def build_docx(blocks: list[Block], out: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for block in blocks:
        if block.kind == "heading":
            doc.add_heading(plain(block.text), level=0 if block.level == 1 else block.level - 1)
        elif block.kind == "paragraph":
            _runs(doc.add_paragraph(), block.text)
        elif block.kind == "bullet":
            _runs(doc.add_paragraph(style="List Bullet"), block.text)
        elif block.kind == "number":
            # Written-out numbers: Word's auto-numbering would run on across lists.
            _runs(doc.add_paragraph(style="List Paragraph"), f"{block.level}. {block.text}")
        elif block.kind == "table":
            table = doc.add_table(rows=0, cols=len(block.rows[0]), style="Table Grid")
            for index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=True):
                    paragraph = cell.paragraphs[0]
                    _runs(paragraph, f"**{plain(text)}**" if index == 0 else text)
    props = doc.core_properties
    props.author = props.last_modified_by = AUTHOR
    props.title = plain(next((b.text for b in blocks if b.kind == "heading"), out.stem))
    props.created = props.modified = props.last_printed = FIXED_DATE.replace(tzinfo=None)
    props.revision = 1
    doc.save(out)
    _fix_zip(out)


# -- XLSX ------------------------------------------------------------------------------


def build_xlsx(blocks: list[Block], out: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    sheet = None
    for block in blocks:
        if block.kind == "heading" and block.level == 2:
            sheet = book.create_sheet(plain(block.text)[:31])
        elif block.kind == "table" and sheet is not None and sheet.max_row == 1:
            widths = [0] * len(block.rows[0])
            for row in block.rows:
                sheet.append([plain(cell) for cell in row])
                widths = [max(w, len(plain(c))) for w, c in zip(widths, row, strict=True)]
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column, width in zip(sheet.columns, widths, strict=True):
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(max(width, 8) + 2, 80)
                for cell in column:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    if not book.sheetnames:
        raise SystemExit(f"{out.name}: no '## Sheet' heading with a table under it")
    book.properties.creator = book.properties.lastModifiedBy = AUTHOR
    book.properties.created = book.properties.modified = FIXED_DATE.replace(tzinfo=None)
    book.save(out)
    _fix_zip(out)


# -- PDF -------------------------------------------------------------------------------


class _Pdf(FPDF):
    def footer(self) -> None:
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")


def _latin1(text: str, source: Path) -> str:
    try:
        text.encode("latin-1")
    except UnicodeEncodeError as exc:
        message = f"{source.name}: the PDF's built-in font is Latin-1 only: {exc}"
        raise SystemExit(message) from None
    return text


def build_pdf(blocks: list[Block], out: Path, source: Path) -> None:
    pdf = _Pdf(format="A4")
    pdf.set_creation_date(FIXED_DATE)
    pdf.set_author(AUTHOR)
    pdf.set_creator(AUTHOR)
    pdf.set_title(plain(next((b.text for b in blocks if b.kind == "heading"), out.stem)))
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    sizes = {1: 18, 2: 14, 3: 12}
    for block in blocks:
        text = _latin1(block.text, source)
        # fpdf2's markdown knows **bold** but not `code`: drop the backticks.
        text = text.replace("`", "")
        if block.kind == "heading":
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", sizes[block.level])
            pdf.multi_cell(0, sizes[block.level] * 0.5, plain(text))
            pdf.ln(1)
        elif block.kind == "paragraph":
            pdf.set_font("Helvetica", size=10)
            pdf.multi_cell(0, 5, text, markdown=True)
            pdf.ln(2)
        elif block.kind in ("bullet", "number"):
            pdf.set_font("Helvetica", size=10)
            mark = "-" if block.kind == "bullet" else f"{block.level}."
            pdf.set_x(pdf.l_margin + 4)
            pdf.cell(6, 5, mark)
            pdf.multi_cell(0, 5, text, markdown=True)
            pdf.ln(1)
        elif block.kind == "table":
            pdf.set_font("Helvetica", size=8)
            heading = FontFace(emphasis="BOLD", fill_color=(230, 230, 230))
            with pdf.table(headings_style=heading, line_height=4.5) as table:
                for row in block.rows:
                    cells = table.row()
                    for cell in row:
                        cells.cell(_latin1(plain(cell), source))
            pdf.ln(3)
    pdf.output(str(out))


# -- deterministic zips ----------------------------------------------------------------

_ZIP_DATE = (2025, 1, 1, 0, 0, 0)
_CORE_DATES = re.compile(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*")


def _fix_zip(path: Path) -> None:
    """Rewrite an Office file with a fixed timestamp on every entry and fixed dates in
    ``docProps/core.xml`` (the libraries stamp the current time on both)."""
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info)) for info in source.infolist()]
    stamp = FIXED_DATE.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = _CORE_DATES.sub(lambda m: m[1] + stamp, data)
            info = zipfile.ZipInfo(name, date_time=_ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            info.create_system = 3  # Unix, whatever the building OS
            target.writestr(info, data)


# -- main ------------------------------------------------------------------------------


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print("usage: uv run samples/build_docs.py samples/<scenario>", file=sys.stderr)
        return 2
    docs = Path(argv[1]) / "docs"
    if not docs.is_dir():
        print(f"{docs} is not a folder", file=sys.stderr)
        return 2
    for source in sorted(docs.glob("*.md")):
        text = source.read_text(encoding="utf-8")
        marker = _MARKER.match(text)
        if marker is None:
            print(f"{source.name}: ships as Markdown")
            continue
        kind = marker[1]
        out = source.with_suffix(f".{kind}")
        blocks = parse(text)
        if kind == "docx":
            build_docx(blocks, out)
        elif kind == "xlsx":
            build_xlsx(blocks, out)
        else:
            build_pdf(blocks, out, source)
        print(f"{source.name} -> {out.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
